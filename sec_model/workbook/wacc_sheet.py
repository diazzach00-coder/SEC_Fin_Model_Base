"""WACC sheet: cost of capital build-up, beta selection and a beta x MRP sensitivity grid.

Mirrors the WACC tab of the analyst's ``DCFWacc Templates.xlsx`` (cost of debt, cost of
equity, capital weights with a target override, WACC; observed-or-industry beta with a
6-row peer table) with every formula live. The share price, unit count and net-debt
bridge also live here (right-hand block) so the DCF sheet links to one source.

Blank-safety: the Model's last history column can hold ``""`` for an unreported concept,
so those links are wrapped in ``N()``; the price input may be blank, so market cap is
``IF(ISNUMBER(price), price*units, 0)``.
"""
from dataclasses import dataclass

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from sec_model.types import MarketInputs
from . import styles as S
from .inputs_sheets import FONT_FORMULA_BOLD, FONT_NOTE, _input, section_header
from .model_sheet import ModelLayout

WACC_SHEET = "WACC"
FONT_CCY_NOTE = Font(name=S.FONT_NAME, size=10, italic=True, color="C00000")
N_PEERS = 6
BASIC_UNAVAILABLE = "Diluted units (Model) — basic unavailable"
STEPS = (-2, -1, 0, 1, 2)          # sensitivity offsets, in units of the grid's step


@dataclass
class WaccLayout:
    wacc_cell: str                 # absolute refs, e.g. "WACC!$E$20"
    tax_cell: str
    kd_cell: str                   # pre-tax cost of debt (input)
    kd_after_tax_cell: str
    ke_cell: str
    beta_cell: str
    equity_weight_cell: str
    debt_weight_cell: str
    price_cell: str                # blank when no price was supplied
    basic_units_cell: str
    diluted_units_cell: str
    market_cap_cell: str
    net_debt_cell: str
    blank_price_equity_weight_cell: str   # We used while the price is blank (input)
    sheet: str = WACC_SHEET


def absref(sheet: str, c: int, r: int) -> str:
    return f"{sheet}!${S.col(c)}${r}"


def offset(base: str, k: int, step: float) -> str:
    """``=base``, ``=base-0.01``, ``=base+0.5`` ... for sensitivity headers."""
    if k == 0:
        return f"={base}"
    return f"={base}{'+' if k > 0 else '-'}{abs(k) * step:g}"


def currency_note(ws, currency: str) -> None:
    if currency != "USD":
        ws.cell(4, 2, f"All figures in {currency} MM. Enter share price in {currency}.").font = \
            FONT_CCY_NOTE


def label(ws, r: int, c: int, text, bold: bool = False):
    cell = ws.cell(r, c, text)
    cell.font = S.FONT_BOLD if bold else S.FONT_BASE
    return cell


def formula(ws, r: int, c: int, value, fmt: str, bold: bool = False, output: bool = False):
    cell = ws.cell(r, c, value)
    cell.font = FONT_FORMULA_BOLD if (bold or output) else S.FONT_FORMULA
    cell.number_format = fmt
    if output:
        cell.fill = S.FILL_OUTPUT
    return cell


def headers(ws, r: int, cols: dict[int, str]) -> None:
    for c, text in cols.items():
        cell = ws.cell(r, c, text)
        cell.font = S.FONT_BOLD
        cell.fill = S.FILL_HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _right_section(ws, r: int, text: str) -> None:
    ws.cell(r, 7, text).font = S.FONT_BOLD
    for c in range(7, 12):
        ws.cell(r, c).fill = S.FILL_SECTION


def _units_and_net_debt(ws, layout: ModelLayout, market: MarketInputs, currency: str,
                        basic_units: float | None) -> dict[str, int]:
    """Right-hand block G:K. Returns rows of price, basic, diluted, market cap, net debt."""
    m, last = layout.sheet, S.col(layout.first_fcst_col - 1)
    lr = layout.line_row

    def model(key):
        return f"N({m}!{last}{lr[key]})"

    _right_section(ws, 5, "SHARE PRICE & UNITS")
    label(ws, 6, 7, f"Share price ({currency} per unit)")
    _input(ws.cell(6, 10), market.price, "perunit")
    if market.price is None:
        ws.cell(6, 11, "enter price").font = FONT_CCY_NOTE
    if basic_units is None:
        label(ws, 7, 7, BASIC_UNAVAILABLE)
        formula(ws, 7, 10, f"={model('UNITS')}", S.FMT_UNITS)
        ws.cell(7, 11, "Already diluted: keep the dilution inputs below at 0").font = FONT_NOTE
    else:
        label(ws, 7, 7, "Basic units (MM)")
        cell = ws.cell(7, 10, basic_units)
        cell.font, cell.number_format = S.FONT_INPUT, S.FMT_UNITS
        ws.cell(7, 11, "dei cover-page units outstanding").font = FONT_NOTE
    for r, text in enumerate(("Restricted units / RSUs", "Options / warrants", "Convertible debt",
                              "Convertible preferred"), start=8):
        label(ws, r, 7, text)
        _input(ws.cell(r, 10), 0, "units")
    label(ws, 12, 7, "Diluted units (MM)", bold=True)
    formula(ws, 12, 10, "=SUM(J7:J11)", S.FMT_UNITS, bold=True)
    label(ws, 13, 7, "Market cap")
    formula(ws, 13, 10, "=IF(ISNUMBER(J6),J6*J12,0)", S.FMT_MONEY)

    _right_section(ws, 15, "NET DEBT (LAST REPORTED YEAR)")
    headers(ws, 16, {8: "Model", 9: "Override", 10: "Used"})
    bridge = (("Debt (incl. revolver)", f"={model('DEBT')}+{model('REVOLVER')}"),
              ("Preferred equity", 0),
              ("Noncontrolling interests", f"={model('NCI_BS')}"),
              ("Less: cash", f"={model('CASH')}"),
              ("Less: equity investments",
               f"=(1-{layout.scalar_cell['JV_TOGGLE']})*{model('EQ_INV')}"))
    for r, (text, value) in enumerate(bridge, start=17):
        label(ws, r, 7, text)
        formula(ws, r, 8, value, S.FMT_MONEY)
        _input(ws.cell(r, 9), None, "money")
        formula(ws, r, 10, f'=IF(I{r}<>"",I{r},H{r})', S.FMT_MONEY)
    ws.cell(21, 11, "JVs valued at book when excluded from EBITDA; via distributions "
                    "when included").font = FONT_NOTE
    label(ws, 22, 7, "Net debt", bold=True)
    formula(ws, 22, 10, "=J17+J18+J19-J20-J21", S.FMT_MONEY, bold=True)
    ws.cell(22, 10).border = S.BORDER_TOP
    return {"price": 6, "basic": 7, "diluted": 12, "mcap": 13, "net_debt": 22}


def write_wacc(wb, layout: ModelLayout, market: MarketInputs, currency: str,
               basic_units: float | None = None) -> WaccLayout:
    """Write the WACC sheet; return its key cells as absolute refs.

    ``basic_units``: latest dei units outstanding (MM); None links the Model's last
    reported diluted units instead."""
    ws = S.get_sheet(wb, WACC_SHEET)
    S.title_bar(ws, 2, "WACC Build-up", 11,
                subtitle=f"All figures in {currency} MM except per-unit data; "
                         "yellow = input, green = result")
    currency_note(ws, currency)
    S.set_widths(ws, {"A": 2, "B": 34, "C": 12, "D": 12, "E": 12, "F": 12, "G": 30,
                      "H": 12, "I": 12, "J": 12, "K": 30})
    right = _units_and_net_debt(ws, layout, market, currency, basic_units)
    fc = S.col(layout.first_fcst_col)
    dr = layout.driver_row

    section_header(ws, 5, "COST OF CAPITAL ASSUMPTIONS", 5)
    label(ws, 6, 2, "Cost of debt (pre-tax)")
    _input(ws.cell(6, 3), f"={layout.drivers_sheet}!{fc}{dr['INT_RATE']}", "pct", formula=True)
    ws.cell(6, 4, "Drivers interest rate, year 1").font = FONT_NOTE
    label(ws, 7, 2, "Tax rate")
    formula(ws, 7, 3, f"={layout.drivers_sheet}!{fc}{dr['TAX_RATE']}", S.FMT_PCT)
    label(ws, 8, 2, "Cost of debt (after tax)")
    formula(ws, 8, 3, "=C6*(1-C7)", S.FMT_PCT, bold=True)
    label(ws, 10, 2, "Risk-free rate")
    _input(ws.cell(10, 3), market.risk_free, "pct")
    label(ws, 11, 2, "Beta")
    formula(ws, 11, 3, '=IF(C33="Observed",C34,C35)', "0.00")
    label(ws, 12, 2, "Market risk premium")
    _input(ws.cell(12, 3), market.mrp, "pct")
    label(ws, 13, 2, "Cost of equity")
    formula(ws, 13, 3, "=C10+C11*C12", S.FMT_PCT, bold=True)

    section_header(ws, 15, "CAPITAL STRUCTURE", 5)
    headers(ws, 16, {3: "Current", 4: "Target (override)", 5: "% of total"})
    label(ws, 17, 2, "Equity (market cap)")
    formula(ws, 17, 3, f"=J{right['mcap']}", S.FMT_MONEY)
    _input(ws.cell(17, 4), None, "pct")
    formula(ws, 17, 5, f'=IF(D17<>"",D17,IF(J{right["price"]}="",D19,'
                       'IFERROR(C17/(C17+C18),0)))', S.FMT_PCT)
    label(ws, 18, 2, "Net debt")
    formula(ws, 18, 3, f"=J{right['net_debt']}", S.FMT_MONEY)
    formula(ws, 18, 5, "=1-E17", S.FMT_PCT)
    label(ws, 19, 2, "Equity weight when price is blank")
    _input(ws.cell(19, 4), 0.55, "pct")
    label(ws, 20, 2, "Cost of capital (WACC)", bold=True)
    formula(ws, 20, 5, "=C8*E18+C13*E17", S.FMT_PCT, output=True)

    section_header(ws, 24, "WACC SENSITIVITY: BETA (COLUMNS) x MARKET RISK PREMIUM (ROWS)", 11)
    headers(ws, 25, {3: "MRP \\ Beta"})
    for j, k in enumerate(STEPS):
        formula(ws, 25, 4 + j, offset("C11", k, 0.1), "0.00", bold=True)
        r = 26 + j
        formula(ws, r, 3, offset("C12", k, 0.01), S.FMT_PCT, bold=True)
        for i in range(len(STEPS)):
            L = S.col(4 + i)
            formula(ws, r, 4 + i, f"=$C$8*$E$18+($C$10+{L}$25*$C{r})*$E$17", S.FMT_PCT,
                    output=(i == 2 and k == 0))

    section_header(ws, 32, "BETA SELECTION", 5)
    label(ws, 33, 2, "Select beta for model")
    _input(ws.cell(33, 3), "Observed", None)
    dv = DataValidation(type="list", formula1='"Observed,Industry"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C33")
    label(ws, 34, 2, "Observed beta")
    _input(ws.cell(34, 3), market.beta, None)
    ws.cell(34, 3).number_format = "0.00"
    label(ws, 35, 2, "Industry beta (relevered)")
    formula(ws, 35, 3, "=C47", "0.00")
    ws.cell(35, 4, "Falls back to the observed beta until peer rows and a price are "
                   "entered").font = FONT_NOTE

    section_header(ws, 37, "INDUSTRY BETA (PEER TABLE)", 9)
    headers(ws, 38, {2: "Peer", 3: "Observed beta", 4: "Price", 5: "Units (MM)",
                     6: "Market cap", 7: "Net debt", 8: "Tax rate", 9: "Unlevered beta"})
    ws.cell(38, 2).alignment = Alignment(horizontal="left")
    for r in range(39, 39 + N_PEERS):
        for c, fmt in ((2, None), (3, None), (4, "perunit"), (5, "units"), (7, "money"),
                       (8, "pct")):
            _input(ws.cell(r, c), None, fmt)
        formula(ws, r, 6, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}*E{r},"")', S.FMT_MONEY)
        formula(ws, r, 9, f'=IFERROR(C{r}/(1+(1-H{r})*G{r}/F{r}),"")', "0.00")
    last_peer = 38 + N_PEERS
    label(ws, 46, 2, "Average unlevered beta")
    formula(ws, 46, 3, f'=IFERROR(AVERAGE(I39:I{last_peer}),"")', "0.00")
    label(ws, 47, 2, "Relevered to subject", bold=True)
    formula(ws, 47, 3, f"=IFERROR(C46*(1+(1-C7)*J{right['net_debt']}/J{right['mcap']}),C34)",
            "0.00", bold=True)

    ws.freeze_panes = "A5"
    a = lambda c, r: absref(WACC_SHEET, c, r)  # noqa: E731
    return WaccLayout(
        wacc_cell=a(5, 20), tax_cell=a(3, 7), kd_cell=a(3, 6), kd_after_tax_cell=a(3, 8),
        ke_cell=a(3, 13), beta_cell=a(3, 11), equity_weight_cell=a(5, 17),
        debt_weight_cell=a(5, 18), price_cell=a(10, right["price"]),
        basic_units_cell=a(10, right["basic"]), diluted_units_cell=a(10, right["diluted"]),
        market_cap_cell=a(10, right["mcap"]), net_debt_cell=a(10, right["net_debt"]),
        blank_price_equity_weight_cell=a(4, 19),
    )
