"""DCF sheet: UFCF build from the Model, discounting, terminal value both ways,
equity bridge, multiples, value-per-share (per-unit) sensitivity grids and a football field.

Mirrors the DCF tab of ``DCFWacc Templates.xlsx``. Sensitivities are ordinary formulas
(no Excel Data Tables): each grid cell recomputes value per share / unit in closed form
from the helper rows ``t`` (years from the valuation date) and stub-adjusted UFCF:
``(SUMPRODUCT(UFCF, 1/(1+w)^t) + TV(w, g|m)/(1+w)^t_N - net debt) / units``.
Price, share / unit count and net debt come from the WACC sheet (one source for both
sheets). Shares vs units follows ``ModelLayout.entity_type`` (`labels`).

JV value is counted once, following ``JV_TOGGLE``: equity earnings are always reversed
out of the EBIT that is taxed and capitalised; with the toggle on, cash distributions
received are added to UFCF; with it off, the JVs enter the net-debt bridge at book.
"""
from dataclasses import dataclass, field
from datetime import date

from openpyxl.chart import BarChart, Reference

from sec_model.labels import label as entity_label, unit_noun
from sec_model.types import MarketInputs
from . import styles as S
from .inputs_sheets import FONT_NOTE, _input, add_toggle_validation, section_header
from .model_sheet import ModelLayout
from .wacc_sheet import (
    STEPS, WaccLayout, absref, currency_note, formula, headers, label, offset,
)

DCF_SHEET = "DCF"
FMT_DATE = "mmm d, yyyy"
FC = 5                              # first forecast column (E); D = last actual year
WACC_C, PRICE_C, MID_C = "$H$6", "$H$7", "$H$8"
VAL_DATE, STUB = "$D$10", "$D$11"


@dataclass
class DcfLayout:
    header_row: int
    row: dict[str, int]                         # block key -> DCF row
    ev_cells: dict[str, str]                    # "perpetuity" | "exit" -> "DCF!$C$45"
    equity_value_cells: dict[str, str]
    value_per_unit_cells: dict[str, str]
    over_under_cells: dict[str, str]            # price / value - 1
    growth_cell: str
    exit_multiple_cell: str
    grids: dict[str, dict] = field(default_factory=dict)   # name -> {"corner": (r, c), "range"}
    grid_centre_cells: dict[str, str] = field(default_factory=dict)
    football_rows: list[int] = field(default_factory=list)  # perpetuity, exit, 52-week
    first_fcst_col: int = FC
    last_fcst_col: int = FC
    sheet: str = DCF_SHEET


# label, key, number format, bold
_UFCF_ROWS = (("Revenue", "REV", S.FMT_MONEY, False),
              ("EBITDA", "EBITDA", S.FMT_MONEY, True),
              ("EBITDA margin", "MARGIN", S.FMT_PCT, False),
              ("EBIT", "EBIT", S.FMT_MONEY, False),
              ("Less: equity earnings (accrual)", "EQ_EARN", S.FMT_MONEY, False),
              ("Tax rate", "TAX_RATE", S.FMT_PCT, False),
              ("Tax on EBIT", "TAX", S.FMT_MONEY, False),
              ("NOPAT", "NOPAT", S.FMT_MONEY, True),
              ("Depreciation & amortization", "DA", S.FMT_MONEY, False),
              ("Change in working capital", "DWC", S.FMT_MONEY, False),
              ("Capital expenditures", "CAPEX", S.FMT_MONEY, False),
              ("JV distributions received", "JV_DIST", S.FMT_MONEY, False),
              ("Unlevered free cash flow (UFCF)", "UFCF", S.FMT_MONEY, True))
_MODEL_LINE = {"REV": "REV", "EBITDA": "EBITDA", "EBIT": "EBIT", "DA": "DA", "DWC": "CF_DWC",
               "CAPEX": "CAPEX"}


def _ufcf_cell(key: str, c: int, row: dict, model_col: str, layout: ModelLayout) -> str:
    L, r = S.col(c), row
    m, lr = layout.sheet, layout.line_row
    actual = c == FC - 1

    def model(line):
        ref = f"{m}!{model_col}{lr[line]}"
        return f"N({ref})" if actual else ref

    if key in _MODEL_LINE:
        return f"={m}!{model_col}{lr[_MODEL_LINE[key]]}"
    if key == "EQ_EARN":
        return f"=-{model('EQ_EARN')}"
    if key == "JV_DIST":
        return f"={layout.scalar_cell['JV_TOGGLE']}*{model('CF_DIST_AFF')}"
    if key == "MARGIN":
        return (f"=IFERROR(N({L}{r['EBITDA']})/N({L}{r['REV']}),0)" if actual
                else f"=IFERROR({L}{r['EBITDA']}/{L}{r['REV']},0)")
    if key == "TAX_RATE":
        if actual:
            return f"=IFERROR({model('TAX')}/{model('PRETAX')},0)"
        return f"={layout.drivers_sheet}!{model_col}{layout.driver_row['TAX_RATE']}"
    ebit = f"N({L}{r['EBIT']})" if actual else f"{L}{r['EBIT']}"
    if key == "TAX":
        return f"=({ebit}+{L}{r['EQ_EARN']})*{L}{r['TAX_RATE']}"
    if key == "NOPAT":
        return f"={ebit}+{L}{r['EQ_EARN']}-{L}{r['TAX']}"
    parts = ([f"{L}{r['NOPAT']}"]
             + [f"N({L}{r[k]})" if actual else f"{L}{r[k]}" for k in ("DA", "DWC", "CAPEX")]
             + [f"{L}{r['JV_DIST']}"])
    return "=" + "+".join(parts)


def _header(ws, wacc: WaccLayout, market: MarketInputs, last_fye: date, company: str,
            ticker: str, entity_type: str) -> None:
    for r, text, value in ((6, "Company name", company), (7, "Ticker", ticker)):
        label(ws, r, 2, text)
        label(ws, r, 4, value)
    label(ws, 8, 2, "Most recent fiscal year end")
    formula(ws, 8, 4, last_fye, FMT_DATE).font = S.FONT_INPUT
    label(ws, 9, 2, "End of first forecast year")
    formula(ws, 9, 4, "=EDATE(D8,12)", FMT_DATE)
    label(ws, 10, 2, "Valuation date")
    formula(ws, 10, 4, "=TODAY()", FMT_DATE)
    label(ws, 11, 2, "Year-1 fraction remaining (stub)")
    formula(ws, 11, 4, "=IF(D10>=D9,0,MIN(1,YEARFRAC(D10,D9)))", "0.000")
    label(ws, 6, 6, "Discount rate (WACC)")
    formula(ws, 6, 8, f"={wacc.wacc_cell}", S.FMT_PCT, bold=True)
    label(ws, 7, 6, f"{unit_noun(entity_type).capitalize()} price")
    p = wacc.price_cell
    formula(ws, 7, 8, f'=IF(ISNUMBER({p}),{p},"")', S.FMT_PERUNIT)
    if market.price is None:
        ws.cell(7, 9, "enter price on WACC sheet").font = FONT_NOTE
    label(ws, 8, 6, "Mid-year convention (1 = on)")
    _input(ws.cell(8, 8), 0, None)
    ws.cell(8, 8).number_format = "0"
    add_toggle_validation(ws, MID_C.replace("$", ""))


def _terminal_value(ws, row: dict, last: str, wacc: WaccLayout, top: int) -> None:
    """Perpetuity method in B:C, exit-multiple method in E:H; ``top`` is the method-title
    row. Adds G, TV, TV_PV, STAGE1, EV_TV and IMPLIED to ``row``."""
    r = row
    g, x1g = top + 1, top + 2
    r.update(G=g, TV=top + 3, TV_PV=top + 4, STAGE1=top + 5, EV_TV=top + 6, IMPLIED=top + 9)
    tv, pv, s1, ev = r["TV"], r["TV_PV"], r["STAGE1"], r["EV_TV"]
    pv_rng = f"{S.col(FC)}{r['PV']}:{last}{r['PV']}"
    df_n = f"{last}{r['DF']}"
    label(ws, top, 2, "Perpetuity growth method", bold=True)
    label(ws, top, 5, "EBITDA exit multiple method", bold=True)
    label(ws, g, 2, "Long-term growth rate (g)")
    _input(ws.cell(g, 3), 0.025, "pct")
    label(ws, g, 5, "Terminal-year EBITDA")
    formula(ws, g, 8, f"={last}{r['EBITDA']}", S.FMT_MONEY)
    label(ws, x1g, 2, "Terminal-year UFCF x (1+g)")
    formula(ws, x1g, 3, f"={last}{r['UFCF']}*(1+C{g})", S.FMT_MONEY)
    label(ws, x1g, 5, "EBITDA exit multiple")
    p = wacc.price_cell
    _input(ws.cell(x1g, 8), f'=IFERROR(IF({p}="",10,({wacc.market_cap_cell}+'
                            f'{wacc.net_debt_cell})/N($D${r["EBITDA"]})),10)', None, formula=True)
    ws.cell(x1g, 8).number_format = S.FMT_X
    ws.cell(x1g, 9, "default: current EV / last-year EBITDA (10.0x without a price)").font = \
        FONT_NOTE
    rows = ((tv, "Terminal value", f"=C{x1g}/($H$6-C{g})", f"=H{g}*H{x1g}", S.FMT_MONEY, False),
            (pv, "PV of terminal value", f"=C{tv}*{df_n}", f"=H{tv}*{df_n}", S.FMT_MONEY, False),
            (s1, "PV of stage-1 cash flows", f"=SUM({pv_rng})", f"=SUM({pv_rng})",
             S.FMT_MONEY, False),
            (ev, "Enterprise value", f"=C{pv}+C{s1}", f"=H{pv}+H{s1}", S.FMT_MONEY, True),
            (ev + 1, "Terminal value % of EV", f"=IFERROR(C{pv}/C{ev},0)",
             f"=IFERROR(H{pv}/H{ev},0)", S.FMT_PCT, False),
            (ev + 2, "Stage-1 cash flows % of EV", f"=IFERROR(C{s1}/C{ev},0)",
             f"=IFERROR(H{s1}/H{ev},0)", S.FMT_PCT, False))
    for rr, text, perp, exit_, fmt, out in rows:
        label(ws, rr, 2, text, bold=out)
        label(ws, rr, 5, text, bold=out)
        formula(ws, rr, 3, perp, fmt, output=out)
        formula(ws, rr, 8, exit_, fmt, output=out)
    im = r["IMPLIED"]
    label(ws, im, 2, "Implied exit EBITDA multiple")
    formula(ws, im, 3, f"=IFERROR(C{tv}/{last}{r['EBITDA']},0)", S.FMT_X)
    label(ws, im, 5, "Implied perpetuity growth rate")
    formula(ws, im, 8, f"=IFERROR(($H$6*H{tv}-{last}{r['UFCF']})/(H{tv}+{last}{r['UFCF']}),0)",
            S.FMT_PCT)


def _grid(ws, top: int, title: str, corner_text: str, col_base: str, col_step: float,
          col_fmt: str, tv: str, row: dict, last: str, wacc: WaccLayout) -> dict:
    """5x5 value-per-share (per-unit) grid; WACC down column C, ``col_base`` offsets across D:H.
    ``tv`` is the terminal-value expression with ``{w}`` and ``{x}`` placeholders."""
    label(ws, top - 1, 2, title, bold=True)
    headers(ws, top, {3: corner_text})
    first = S.col(FC)
    ufcf = f"${first}${row['UFCF_ADJ']}:${last}${row['UFCF_ADJ']}"
    t = f"${first}${row['T']}:${last}${row['T']}"
    t_n = f"${last}${row['T']}"
    for j, k in enumerate(STEPS):
        formula(ws, top, 4 + j, offset(col_base, k, col_step), col_fmt, bold=True)
        r = top + 1 + j
        formula(ws, r, 3, offset(WACC_C, k, 0.005), S.FMT_PCT, bold=True)
        for i in range(len(STEPS)):
            w, x = f"$C{r}", f"{S.col(4 + i)}${top}"
            f = (f"=IFERROR((SUMPRODUCT({ufcf},1/(1+{w})^{t})+{tv.format(w=w, x=x)}"
                 f"/(1+{w})^{t_n}-{wacc.net_debt_cell})/{wacc.diluted_units_cell},\"n/a\")")
            formula(ws, r, 4 + i, f, S.FMT_PERUNIT, output=(i == 2 and k == 0))
    return {"corner": (top, 3), "range": f"D{top + 1}:H{top + 5}"}


def _value_per(entity_type: str) -> str:
    """"Value per share" / "Value per unit"."""
    return f"Value {entity_label('per_unit', entity_type)}"


def football_chart(ws, rows: list[int], anchor: str, title: str) -> None:
    """Stacked horizontal bar: invisible 'Low' series + visible 'Spread' series."""
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "bar", "stacked", 100
    chart.title = title
    chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=rows[0] - 1, max_row=rows[-1]),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=rows[0], max_row=rows[-1]))
    chart.series[0].graphicalProperties.noFill = True
    chart.series[0].graphicalProperties.line.noFill = True
    chart.legend = None
    chart.height, chart.width = 6.5, 16
    ws.add_chart(chart, anchor)


def _ufcf_block(ws, layout: ModelLayout, fcst_labels: list[str], lastc: int) -> dict:
    """Rows 13-28: header, fiscal-year-end dates and the UFCF build. Returns ``row``."""
    n = len(fcst_labels)
    section_header(ws, 13, "UNLEVERED FREE CASH FLOW", lastc)
    hist_label = ws.parent[layout.sheet].cell(4, layout.first_fcst_col - 1).value or "Actual"
    S.header_row(ws, 14, ["Fiscal year"], 2)
    S.header_row(ws, 14, [hist_label] + list(fcst_labels), 4, [False] + [True] * n)
    row = {"FYE": 15}
    label(ws, 15, 2, "Fiscal year ended")
    formula(ws, 15, 4, "=D8", FMT_DATE)
    for c in range(FC, lastc + 1):
        formula(ws, 15, c, f"=EDATE({S.col(c - 1)}15,12)", FMT_DATE)
    for i, (_, key, _, _) in enumerate(_UFCF_ROWS):
        row[key] = 16 + i
    for text, key, fmt, bold in _UFCF_ROWS:
        r = row[key]
        label(ws, r, 2, text, bold=bold)
        for c in range(FC - 1, lastc + 1):
            model_col = S.col(layout.first_fcst_col + c - FC)   # DCF D = last Model history
            cell = formula(ws, r, c, _ufcf_cell(key, c, row, model_col, layout), fmt, bold=bold)
            if c >= FC:
                cell.fill = S.FILL_FORECAST
            if key == "UFCF":
                cell.border = S.BORDER_TOP
    return row


def _discounting(ws, row: dict, n: int, lastc: int) -> None:
    top = row["UFCF"] + 2
    section_header(ws, top, "PRESENT VALUE OF UFCF", lastc)
    headers(ws, top + 1, {4: "Valuation date", **{FC + i: ("Year 1 (stub)" if i == 0 else
                                                           f"Year {i + 1}") for i in range(n)}})
    dd, t, adj, df, pv = range(top + 2, top + 7)
    row.update(DISC_DATE=dd, T=t, UFCF_ADJ=adj, DF=df, PV=pv)
    for r, text in ((dd, "Discount date"), (t, "Years from valuation date (t)"),
                    (adj, "UFCF (year 1 stub-adjusted)"), (df, "Discount factor"),
                    (pv, "PV of UFCF")):
        label(ws, r, 2, text, bold=(r == pv))
    formula(ws, dd, 4, "=D10", FMT_DATE)
    for c in range(FC, lastc + 1):
        L = S.col(c)
        prev = VAL_DATE if c == FC else f"{S.col(c - 1)}15"
        formula(ws, dd, c, f"=IF({MID_C}=0,{L}15,AVERAGE({prev},{L}15))", FMT_DATE)
        formula(ws, t, c, f"=({L}{dd}-{VAL_DATE})/365", "0.000")
        formula(ws, adj, c, f"={L}{row['UFCF']}*{STUB}" if c == FC else f"={L}{row['UFCF']}",
                S.FMT_MONEY)
        formula(ws, df, c, f"=1/(1+{WACC_C})^{L}{t}", "0.0000")
        formula(ws, pv, c, f"={L}{adj}*{L}{df}", S.FMT_MONEY, bold=True)


def _equity_and_multiples(ws, row: dict, wacc: WaccLayout, last: str, lastc: int,
                          entity_type: str) -> int:
    """Equity bridge per method, then implied multiples. Returns the last row written."""
    top = row["IMPLIED"] + 2
    vpu_label = _value_per(entity_type)
    section_header(ws, top, f"EQUITY VALUE & {vpu_label.upper()}", lastc)
    headers(ws, top + 1, {3: "Perpetuity", 4: "Exit multiple"})
    ev, nd, eq, un, vpu, pr, ou = range(top + 2, top + 9)
    row.update(EV=ev, NET_DEBT=nd, EQUITY=eq, UNITS=un, VPU=vpu, PRICE=pr, OVER_UNDER=ou)
    p, evtv = wacc.price_cell, row["EV_TV"]
    for r, text, perp, exit_, fmt, out in (
            (ev, "Enterprise value", f"=C{evtv}", f"=H{evtv}", S.FMT_MONEY, False),
            (nd, "Less: net debt", f"={wacc.net_debt_cell}", f"={wacc.net_debt_cell}",
             S.FMT_MONEY, False),
            (eq, "Equity value", f"=C{ev}-C{nd}", f"=D{ev}-D{nd}", S.FMT_MONEY, False),
            (un, f"{entity_label('units', entity_type)} (MM)", f"={wacc.diluted_units_cell}",
             f"={wacc.diluted_units_cell}", S.FMT_UNITS, False),
            (vpu, vpu_label, f'=IFERROR(C{eq}/C{un},"n/a")',
             f'=IFERROR(D{eq}/D{un},"n/a")', S.FMT_PERUNIT, True),
            (pr, "Current price", f"={PRICE_C}", f"={PRICE_C}", S.FMT_PERUNIT, False),
            (ou, "% over / (under) valued",
             f'=IF({p}="","enter price",IFERROR({p}/C{vpu}-1,"n/a"))',
             f'=IF({p}="","enter price",IFERROR({p}/D{vpu}-1,"n/a"))', S.FMT_PCT, False)):
        label(ws, r, 2, text, bold=out)
        formula(ws, r, 3, perp, fmt, output=out)
        formula(ws, r, 4, exit_, fmt, output=out)

    top = ou + 2
    section_header(ws, top, "IMPLIED MULTIPLES", lastc)
    headers(ws, top + 1, {3: "Current", 4: "Current", 5: "Year 1", 6: "Year 1",
                          7: "Terminal year", 8: "Terminal year"})
    headers(ws, top + 2, {c: ("Perpetuity" if c % 2 else "Exit") for c in range(3, 9)})
    row.update(MULT_REV=top + 3, MULT_EBITDA=top + 4, MULT_EBIT=top + 5)
    for r, text, key in ((top + 3, "EV / Revenue", "REV"), (top + 4, "EV / EBITDA", "EBITDA"),
                         (top + 5, "EV / EBIT", "EBIT")):
        label(ws, r, 2, text)
        rr = row[key]
        for c, evc, denom in ((3, f"$C${ev}", f"N($D${rr})"), (4, f"$D${ev}", f"N($D${rr})"),
                              (5, f"$C${ev}", f"E{rr}"), (6, f"$D${ev}", f"E{rr}"),
                              (7, f"$C${ev}", f"{last}{rr}"), (8, f"$D${ev}", f"{last}{rr}")):
            formula(ws, r, c, f'=IFERROR({evc}/{denom},"n/a")', S.FMT_X)
    ws.cell(top + 5, 9, "EBIT includes equity-method earnings (reported basis)").font = FONT_NOTE
    return top + 5


def _football_field(ws, top: int, grids: dict, lastc: int, entity_type: str) -> list[int]:
    section_header(ws, top, f"FOOTBALL FIELD: {_value_per(entity_type).upper()}", lastc)
    headers(ws, top + 1, {3: "Low", 4: "Spread", 5: "High"})
    ff = [top + 2, top + 3, top + 4]
    for r, name, title, fmt in ((ff[0], "perpetuity", "DCF perpetuity: g ", "0.0%"),
                                (ff[1], "exit", "DCF exit multiple: ", "0.0x")):
        gt = grids[name]["corner"][0]
        label(ws, r, 2, f'="{title}"&TEXT($D${gt},"{fmt}")&" to "&TEXT($H${gt},"{fmt}")'
                        f'&", WACC "&TEXT($C${gt + 1},"0.0%")&" to "&TEXT($C${gt + 5},"0.0%")')
        rng = grids[name]["range"]
        formula(ws, r, 3, f"=MIN({rng})", S.FMT_PERUNIT)
        formula(ws, r, 4, f"=E{r}-C{r}", S.FMT_PERUNIT)
        formula(ws, r, 5, f"=MAX({rng})", S.FMT_PERUNIT)
    wk = ff[2]
    label(ws, wk, 2, "52-week trading range (enter low / high)")
    _input(ws.cell(wk, 3), None, "perunit")
    formula(ws, wk, 4, f"=IF(AND(ISNUMBER(C{wk}),ISNUMBER(E{wk})),E{wk}-C{wk},0)",
            S.FMT_PERUNIT)
    _input(ws.cell(wk, 5), None, "perunit")
    return ff


def write_dcf(wb, layout: ModelLayout, wacc: WaccLayout, market: MarketInputs,
              fcst_labels: list[str], currency: str, last_fye: date, *,
              company: str = "", ticker: str = "") -> DcfLayout:
    """Write the DCF sheet off the Model, Drivers and WACC sheets.

    Rows below the UFCF block are placed relative to it; ``DcfLayout.row`` has them all."""
    n = layout.last_col - layout.first_fcst_col + 1
    if n < 1 or n != len(fcst_labels):
        raise ValueError(f"{len(fcst_labels)} forecast labels for {n} Model forecast columns")
    ws = S.get_sheet(wb, DCF_SHEET)
    e = layout.entity_type
    lastc = FC + n - 1
    last = S.col(lastc)
    bar_end = max(lastc, 9)
    S.title_bar(ws, 2, "Discounted Cash Flow Valuation", bar_end,
                subtitle=f"All figures in {currency} MM except per-{unit_noun(e)} data; "
                         "yellow = input, green = result")
    currency_note(ws, currency, e)
    S.set_widths(ws, {"A": 2, "B": 40, "C": 13})
    for c in range(4, bar_end + 1):
        S.set_widths(ws, {c: 13})
    _header(ws, wacc, market, last_fye, company, ticker, e)
    row = _ufcf_block(ws, layout, fcst_labels, lastc)
    _discounting(ws, row, n, lastc)

    top = row["PV"] + 2
    section_header(ws, top, "TERMINAL VALUE", lastc)
    _terminal_value(ws, row, last, wacc, top + 1)
    top = _equity_and_multiples(ws, row, wacc, last, lastc, e) + 2

    section_header(ws, top, f"SENSITIVITY: {_value_per(e).upper()}", lastc)
    last_ufcf, last_ebitda = f"${last}${row['UFCF']}", f"${last}${row['EBITDA']}"
    grids = {
        "perpetuity": _grid(ws, top + 2, "WACC (rows) x long-term growth g (columns)",
                            "WACC \\ g", f"$C${row['G']}", 0.005, S.FMT_PCT,
                            last_ufcf + "*(1+{x})/({w}-{x})", row, last, wacc),
        "exit": _grid(ws, top + 10, "WACC (rows) x exit EBITDA multiple (columns)",
                      "WACC \\ multiple", f"$H${row['G'] + 1}", 0.5, S.FMT_X,
                      last_ebitda + "*{x}", row, last, wacc),
    }
    ff = _football_field(ws, top + 17, grids, lastc, e)
    football_chart(ws, ff, f"B{ff[-1] + 2}", _value_per(e))

    ws.freeze_panes = "C5"
    a = lambda c, r: absref(DCF_SHEET, c, r)  # noqa: E731
    return DcfLayout(
        header_row=14, row=row,
        ev_cells={"perpetuity": a(3, row["EV_TV"]), "exit": a(8, row["EV_TV"])},
        equity_value_cells={"perpetuity": a(3, row["EQUITY"]), "exit": a(4, row["EQUITY"])},
        value_per_unit_cells={"perpetuity": a(3, row["VPU"]), "exit": a(4, row["VPU"])},
        over_under_cells={"perpetuity": a(3, row["OVER_UNDER"]),
                          "exit": a(4, row["OVER_UNDER"])},
        growth_cell=a(3, row["G"]), exit_multiple_cell=a(8, row["G"] + 1), grids=grids,
        grid_centre_cells={k: a(6, g["corner"][0] + 3) for k, g in grids.items()},
        football_rows=ff, first_fcst_col=FC, last_fcst_col=lastc,
    )
