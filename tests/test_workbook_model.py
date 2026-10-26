from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from sec_model.concepts import CONCEPTS
from sec_model.forecast import driver_defaults, scalar_defaults
from sec_model.formula import Const, Expr, to_excel
from sec_model.model_spec import build_spec
from sec_model.workbook import styles as S
from sec_model.workbook.model_sheet import ModelLayout, render_context, write_model
from test_forecast import FCST, FY, TOY

# History D:F (FY2023-FY2025), forecast G:K (FY2026E-FY2030E).
G, H, F = 7, 8, 6


def _mapping_stub(wb, history):
    """A Mapping sheet with literal values in the Task 10 layout (rows from 5, values from D)."""
    ws = S.get_sheet(wb, "Mapping")
    rows = {c.name: 5 + i for i, c in enumerate(CONCEPTS)}
    value_cols = {fy: S.DSC + i for i, fy in enumerate(FY)}
    for name, r in rows.items():
        for fy, c in value_cols.items():
            v = (history.get(name) or {}).get(fy)
            if v is not None:
                ws.cell(r, c, v)
    return rows, value_cols


def _build(history=TOY):
    spec = build_spec("ppe")
    wb = Workbook()
    rows, value_cols = _mapping_stub(wb, history)
    dd = driver_defaults(spec, history, FY)
    sd = scalar_defaults(spec, history, FY)
    layout = write_model(wb, spec, rows, value_cols, FY, FCST, dd, sd, "USD")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf), layout, spec, rows, dd, sd


def _cell(ws, key, c, layout):
    return ws.cell(layout.line_row[key], c)


def test_layout_shape():
    wb, layout, spec, *_ = _build()
    assert isinstance(layout, ModelLayout)
    assert wb.sheetnames == ["Mapping", "Model", "Asset Inputs", "Drivers"]
    assert (layout.sheet, layout.drivers_sheet, layout.assets_sheet) == \
        ("Model", "Drivers", "Asset Inputs")
    assert (layout.first_hist_col, layout.first_fcst_col, layout.last_col) == (4, G, 11)
    assert set(layout.line_row) == set(spec.line_keys())
    assert set(layout.driver_row) == {d.key for d in spec.drivers}
    assert set(layout.scalar_cell) == {s.key for s in spec.scalars}
    ws = wb["Model"]
    assert ws.freeze_panes == "D5"
    assert ws["B2"].value == "3-Statement Model"
    assert [ws.cell(4, c).value for c in range(4, 12)] == FY + FCST
    assert ws.cell(4, G).fill.start_color.rgb.endswith("EBF3FB")


def test_model_history_cell_links_to_mapping():
    wb, layout, spec, rows, *_ = _build()
    ws = wb["Model"]
    for i, fy in enumerate(FY):
        c = 4 + i
        cell = _cell(ws, "REV", c, layout)
        assert cell.value == f"=Mapping!{S.col(c)}{rows['Revenue']}"
        assert cell.font.color.rgb.endswith(S.BLUE)
    # derived history line: rendered with the writer's context
    c = F
    ae = _cell(ws, "ASSET_EBITDA", c, layout)
    assert ae.value == to_excel(spec.line("ASSET_EBITDA").hist, render_context(layout, spec, c))
    assert ae.value == f"=N(F{layout.line_row['REV']})-F{layout.line_row['OPCOST']}"
    assert ae.font.color.rgb.endswith("000000")
    # hist None -> blank, never 0
    for c in range(4, G):
        assert _cell(ws, "CASH_PRE", c, layout).value is None
    assert _cell(ws, "CASH_PRE", G, layout).value is not None
    # concept lines referenced from formulas are blank-safe (Mapping may hold "")
    opcost = _cell(ws, "OPCOST", 4, layout).value
    assert f"N(D{layout.line_row['TCE_RPT']})" in opcost


def test_forecast_cell_formula_matches_renderer():
    wb, layout, spec, *_ = _build()
    ws = wb["Model"]
    r = layout.line_row
    for c in range(G, layout.last_col + 1):
        ctx = render_context(layout, spec, c)
        for ld in spec.lines:
            got = ws.cell(r[ld.key], c).value
            if ld.fcst is None:
                assert got is None, ld.key
            else:
                assert got == to_excel(ld.fcst, ctx), (ld.key, c)
            assert ws.cell(r[ld.key], c).fill.start_color.rgb.endswith("EBF3FB")
    jv = layout.scalar_cell["JV_TOGGLE"]
    assert jv.startswith("Drivers!$C$")
    assert ws.cell(r["EBITDA"], G).value == f"=G{r['ASSET_EBITDA']}+(G{r['EQ_EARN']}*{jv})"
    # first forecast year reads last history cash (blank-safe), later years plain refs
    assert ws.cell(r["CASH"], G).value == f"=N(F{r['CASH']})+G{r['NET_CHANGE']}"
    assert ws.cell(r["CASH"], H).value == f"=G{r['CASH']}+H{r['NET_CHANGE']}"
    assert ws.cell(r["BAL_CHECK"], H).value == f"=H{r['TOTAL_ASSETS']}-H{r['TOTAL_LE']}"
    dr = layout.driver_row
    assert ws.cell(r["INTEREST"], H).value == \
        f"=(Drivers!H{dr['INT_RATE']}*G{r['DEBT']})+(Drivers!H{dr['REV_RATE']}*G{r['REVOLVER']})"


def test_driver_first_col_value_and_chain():
    wb, layout, spec, _, dd, sd = _build()
    ws = wb["Drivers"]
    assert ws.freeze_panes == "D5"
    assert [ws.cell(4, c).value for c in range(4, 12)] == FY + FCST
    assert [ws.cell(layout.driver_row[d.key], 2).value for d in spec.drivers] == \
        [d.label for d in spec.drivers]
    assert layout.driver_row[spec.drivers[0].key] == 5
    ctx = render_context(layout, spec, F, qualified=True)     # last history column
    for d in spec.drivers:
        r = layout.driver_row[d.key]
        first = ws.cell(r, G)
        if isinstance(d.default, Const):                       # OTHER_OP, DEBT_ISSUE
            assert first.value == d.default.value == dd[d.key]
        else:                                                  # spec 7.2: a formula
            assert first.value == to_excel(d.default, ctx), d.key
        assert first.font.color.rgb.endswith(S.BLUE)
        assert first.fill.start_color.rgb.endswith("FFF2CC")
        for c in range(H, 12):
            nxt = ws.cell(r, c)
            assert nxt.value == f"={S.col(c - 1)}{r}"
            assert nxt.fill.start_color.rgb.endswith("FFF2CC")
        assert all(ws.cell(r, c).value is None for c in range(4, G))
        assert ws.cell(r, 12).value == d.note
    for s in spec.scalars:
        ref = layout.scalar_cell[s.key]
        row = int(ref.rsplit("$", 1)[1])
        assert ref == f"Drivers!$C${row}"
        assert ws.cell(row, 2).value == s.label
        if isinstance(s.default, Expr):
            assert ws.cell(row, 3).value == to_excel(s.default, ctx)
        else:
            assert ws.cell(row, 3).value == pytest.approx(sd[s.key], rel=1e-12)
        assert ws.cell(row, 3).fill.start_color.rgb.endswith("FFF2CC")
        assert ws.cell(row, 4).value == s.note
        assert wb.defined_names[s.key].attr_text == ref
    assert set(sd) == {"JV_TOGGLE", "WC_HOLD", "REV_CAP", "GROWTH_CAP"}
    capex = ws.cell(layout.driver_row["CAPEX_PCT"], G).value
    assert capex == f"=AVERAGE(Model!$D${layout.line_row['CAPEX_PCT_HIST']}:$F$"         f"{layout.line_row['CAPEX_PCT_HIST']})"
    assert "AVERAGE(" in capex and "Model!" in capex
    assert ws.cell(layout.driver_row["UNITS"], G).value == f"=N(Model!F{layout.line_row['UNITS']})"
    # per-year minimum cash and revolver rate: year 1 a formula over the Model, then =prev
    r_min, r_rate = layout.driver_row["MIN_CASH"], layout.driver_row["REV_RATE"]
    assert r_rate == r_min + 1 == layout.driver_row["DEBT_ISSUE"] + 2
    assert ws.cell(r_min, 2).value == "Minimum cash balance"
    assert ws.cell(r_min, G).value == f"=N(Model!F{layout.line_row['CASH']})"
    assert ws.cell(r_rate, G).value == f"=Model!F{layout.line_row['INT_RATE_HIST']}"
    assert ws.cell(r_rate, 12).value == "revolver rate; defaults to the term-debt rate"
    for r in (r_min, r_rate):
        assert ws.cell(r, H).value == f"=G{r}"
    assert ws.cell(r_min, G).number_format == S.FMT_MONEY
    assert ws.cell(r_rate, G).number_format == S.FMT_PCT
    # scalars: revolver commitment and growth cap are plain yellow numeric inputs
    for key, val, fmt in (("REV_CAP", 0, S.FMT_MONEY), ("GROWTH_CAP", 0.15, S.FMT_PCT)):
        row = int(layout.scalar_cell[key].rsplit("$", 1)[1])
        assert ws.cell(row, 3).value == val and ws.cell(row, 3).number_format == fmt
        assert ws.cell(row, 3).font.color.rgb.endswith(S.BLUE)
    cap = layout.scalar_cell["GROWTH_CAP"]
    assert ws.cell(layout.driver_row["REV_GROWTH"], G).value ==         f"=MAX(MIN(AVERAGE(Model!$D${layout.line_row['REV_GROWTH_HIST']}:$F$"         f"{layout.line_row['REV_GROWTH_HIST']}),{cap}),-{cap})"


def test_asset_block_total_feeds_revenue():
    wb, layout, spec, *_ = _build()
    ws, wa = wb["Model"], wb["Asset Inputs"]
    r, dr = layout.line_row, layout.driver_row
    rev_tot = layout.asset_total_cell[("ASSET_REV", "FY2026E")]
    cost_tot = layout.asset_total_cell[("ASSET_COST", "FY2027E")]
    assert rev_tot.startswith("'Asset Inputs'!G")
    assert ws.cell(r["REV"], G).value == "=" + rev_tot
    assert ws.cell(r["OPCOST"], H).value == "=" + cost_tot
    tot_row = int(rev_tot.rsplit("G", 1)[1])
    assert wa.cell(tot_row, 2).value == "Total revenue"
    first = tot_row - 25
    assert wa.cell(tot_row, G).value == f"=SUM(G{first}:G{tot_row - 1})"
    assert wa.cell(first, 2).value == "Consolidated"
    assert wa.cell(tot_row - 1, 2).value == "Other / unallocated"
    # placeholder asset: last reported revenue grown by the (per-year) REV_GROWTH driver
    assert wa.cell(first, G).value == \
        f"=N(Model!F{r['REV']})*(1+Drivers!G{dr['REV_GROWTH']})"
    assert wa.cell(first, H).value == f"=G{first}*(1+Drivers!H{dr['REV_GROWTH']})"
    for rr in range(first + 1, tot_row):
        assert wa.cell(rr, G).value is None
        assert wa.cell(rr, G).fill.start_color.rgb.endswith("FFF2CC")
    cost_row = int(cost_tot.rsplit("H", 1)[1])
    cfirst = cost_row - 25
    assert wa.cell(cost_row, 2).value == "Total operating cost"
    assert wa.cell(cfirst, H).value == f"=H{first}*Drivers!H{dr['COST_PCT']}"
    labels = [wa.cell(rr, 2).value for rr in range(cost_row + 1, wa.max_row + 1)]
    rec = labels.index("Last reported revenue (Model)") + cost_row + 1
    assert wa.cell(rec, F).value == f"=Model!F{r['REV']}"
    assert wa.cell(rec, G).value is None
    assert wa.cell(rec + 1, F).value == f"=Model!F{r['OPCOST']}"
    assert wa.cell(rec + 2, 2).value == "Asset EBITDA"
    assert wa.cell(rec + 2, G).value == f"=G{tot_row}-G{cfirst + 25}"


def test_balance_check_has_conditional_format():
    wb, layout, *_ = _build()
    ws = wb["Model"]
    row = layout.line_row["BAL_CHECK"]
    rules = [(str(rng.sqref), rule) for rng in ws.conditional_formatting for rule in rng.rules]
    found = {(rule.operator, tuple(rule.formula)) for sq, rule in rules if sq == f"D{row}:K{row}"}
    assert found == {("greaterThan", ("0.5",)), ("lessThan", ("-0.5",))}
    for sq, rule in rules:
        assert rule.dxf.fill.fgColor.rgb.endswith("F4CCCC") or \
            rule.dxf.fill.bgColor.rgb.endswith("F4CCCC")


def test_liquidity_shortfall_row_and_conditional_format():
    wb, layout, spec, *_ = _build()
    ws = wb["Model"]
    r = layout.line_row
    assert r["LIQ_SHORTFALL"] == r["BAL_CHECK"] + 1
    assert ws.cell(r["LIQ_SHORTFALL"], 2).value == "Liquidity shortfall vs minimum cash"
    assert all(ws.cell(r["LIQ_SHORTFALL"], c).value is None for c in range(4, G))
    assert ws.cell(r["LIQ_SHORTFALL"], H).value ==         f"=MAX(0,Drivers!H{layout.driver_row['MIN_CASH']}-H{r['CASH']})"
    rules = [(str(rng.sqref), rule) for rng in ws.conditional_formatting for rule in rng.rules]
    row = r["LIQ_SHORTFALL"]
    found = [(rule.operator, tuple(rule.formula)) for sq, rule in rules if sq == f"G{row}:K{row}"]
    assert found == [("greaterThan", ("0.5",))]
    # the revolver cell turns red in a year whose need exceeds the commitment
    rev = r["REVOLVER"]
    found = [rule for sq, rule in rules if sq == f"G{rev}:K{rev}"]
    assert len(found) == 1 and found[0].type == "expression"
    assert found[0].formula == [f"G{row}>0.5"]
    for sq, rule in rules:
        assert rule.dxf.fill.fgColor.rgb.endswith("F4CCCC") or             rule.dxf.fill.bgColor.rgb.endswith("F4CCCC")


def test_debt_history_splits_out_the_drawn_revolver():
    h = dict(TOY, RevolverDrawn={"FY2023": 10, "FY2024": 20, "FY2025": 30})
    wb, layout, spec, rows, *_ = _build(h)
    ws = wb["Model"]
    r = layout.line_row
    assert ws.cell(r["REV_RPT"], F).value == f"=Mapping!F{rows['RevolverDrawn']}"
    assert ws.row_dimensions[r["REV_RPT"]].hidden
    assert ws.cell(r["REVOLVER"], F).value == f"=N(F{r['REV_RPT']})"
    assert ws.cell(r["DEBT"], F).value == f"=N(F{r['DEBT_RPT']})-N(F{r['REV_RPT']})"
    assert ws.cell(r["DEBT_NET"], F).value ==         f"=(F{r['DEBT']}+F{r['REVOLVER']})-(E{r['DEBT']}+E{r['REVOLVER']})"


def test_sections_hidden_rows_and_styles():
    wb, layout, spec, *_ = _build()
    ws = wb["Model"]
    r = layout.line_row
    headers = [(ws.cell(rr, 2).value, rr) for rr in range(5, ws.max_row + 1)
               if ws.cell(rr, 2).fill.start_color.rgb.endswith("F2F2F2")]
    assert [h for h, _ in headers] == ["INCOME STATEMENT", "RATIOS & DRIVER HISTORY",
                                       "CASH FLOW STATEMENT", "BALANCE SHEET", "BALANCE CHECK"]
    hrow = dict(headers)
    assert hrow["INCOME STATEMENT"] < r["REV"] < r["EPS"] < hrow["RATIOS & DRIVER HISTORY"] \
        < r["REV_GROWTH_HIST"] < r["DIST_COVERAGE"] < hrow["CASH FLOW STATEMENT"] < r["CF_NI"] \
        < hrow["BALANCE SHEET"] < r["CASH"] < r["TOTAL_LE"] < hrow["BALANCE CHECK"] \
        < r["BAL_CHECK"]
    assert ws.cell(hrow["BALANCE SHEET"], 2).fill.start_color.rgb.endswith("F2F2F2")
    for ld in spec.lines:
        dim = ws.row_dimensions[r[ld.key]]
        assert bool(dim.hidden) == ld.hidden, ld.key
        if ld.hidden:
            assert dim.outlineLevel == 1
    ebitda = ws.cell(r["EBITDA"], 5)
    assert ebitda.font.bold and ebitda.fill.start_color.rgb.endswith("E8E8E8")
    assert ebitda.border.top.style == "thin"
    assert ws.cell(r["EBITDA_MARGIN"], 5).number_format == S.FMT_PCT
    assert ws.cell(r["NET_DEBT_EBITDA"], 5).number_format == S.FMT_X
    assert ws.cell(r["EPS"], 3).value == "per unit"
    assert ws.cell(r["REV"], 3).value == "$MM"
    note_col = layout.last_col + 2
    assert ws.cell(r["EQ_EARN"], note_col).value == "held flat"
    assert ws.cell(r["EQ_EARN"], note_col).font.italic


def test_missing_history_concept_is_blank_not_zero():
    hist = {k: dict(v) for k, v in TOY.items() if k != "Revenue"}
    spec = build_spec("ppe")
    wb = Workbook()
    rows = {c.name: 5 + i for i, c in enumerate(CONCEPTS) if c.name != "Revenue"}
    value_cols = {fy: S.DSC + i for i, fy in enumerate(FY)}
    layout = write_model(wb, spec, rows, value_cols, FY, FCST, driver_defaults(spec, hist, FY),
                         scalar_defaults(spec, hist, FY), "CAD")
    ws = wb["Model"]
    for c in range(4, G):
        assert ws.cell(layout.line_row["REV"], c).value is None
    assert ws.cell(layout.line_row["REV"], 3).value == "CAD MM"


def test_note_starting_with_equals_is_text_not_formula():
    import dataclasses
    spec = build_spec("ppe")
    i = next(i for i, ld in enumerate(spec.lines) if ld.key == "REV")
    spec.lines[i] = dataclasses.replace(spec.lines[i], note="= revenue note, not a formula")
    wb = Workbook()
    rows, value_cols = _mapping_stub(wb, TOY)
    layout = write_model(wb, spec, rows, value_cols, FY, FCST, driver_defaults(spec, TOY, FY),
                         scalar_defaults(spec, TOY, FY), "USD")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    cell = load_workbook(buf)["Model"].cell(layout.line_row["REV"], layout.last_col + 2)
    assert cell.value == "= revenue note, not a formula"
    assert cell.data_type == "s"
