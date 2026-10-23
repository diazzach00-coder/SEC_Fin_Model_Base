"""Regression run: build the workbook for each ticker against live SEC data and assert
the acceptance rules (spec section 10).

    python scripts/regress.py --ua "Name email@example.com" OKE EPD ET ... TRP

Per ticker: writes ``<out>/<TICKER>.xlsx`` and prints one line (years found, statements
per year, check counts, coverage warnings, max |BAL_CHECK| over the forecast years,
elapsed). Assertion failures are collected; the exit code is 1 if any ticker failed.

Assertions: no exception, except a ``UserFacingError`` mentioning ``40-F`` for the 40-F
filers in ``EXPECT_40F``; at least ``--min-years`` fiscal years; IS, BS and CF for every
kept year; zero FAIL checks; max |BAL_CHECK| < 1e-6; the file reopens with openpyxl;
no ``#REF!`` in any cell.

Sanity notes (printed as SANITY lines; assertions only with ``--strict``): the default
interest rate on debt is negative, or an expense-type concept's latest mapped value has the
opposite sign to the companyfacts value for the same tag and year (an R-page "negated
label" carried into the model), or interest expense / pretax income is unmapped.
"""
import argparse
import logging
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from sec_model.checks import summarize  # noqa: E402
from sec_model.companyfacts import fy_values, load_companyfacts  # noqa: E402
from sec_model.http import SecClient  # noqa: E402
from sec_model.pipeline import STYPES, build  # noqa: E402
from sec_model.types import UserFacingError  # noqa: E402

EXPECT_40F = {"TRP"}
BAL_TOL = 1e-6
# Model convention: these are positive when they reduce income (NIToNCI: positive = NCI's
# share of income), which is also the XBRL sign of their us-gaap tags.
SIGN_CONCEPTS = ("InterestExpense", "IncomeTax", "NIToNCI", "DA")


@dataclass
class Row:
    ticker: str
    years: str = ""
    n_years: int = 0
    stmts: str = ""
    counts: dict = field(default_factory=dict)
    coverage: list = field(default_factory=list)
    bal: float | None = None
    secs: float = 0.0
    kb: int = 0
    currency: str = ""
    note: str = ""
    failures: list = field(default_factory=list)
    sanity: list = field(default_factory=list)

    @property
    def status(self) -> str:
        return "FAIL" if self.failures else "ok"


def scan_ref_errors(path: Path) -> tuple[list[str], str | None]:
    """Cells holding '#REF!' (value or formula text), and an openpyxl load error if any."""
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:  # any load failure is an assertion failure, not a crash
        return [], f"{type(e).__name__}: {e}"
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "#REF!" in c.value:
                    hits.append(f"{ws.title}!{c.coordinate}")
    return hits, None


def sanity_notes(res, ua: str) -> list[str]:
    """Model-semantics red flags that the acceptance assertions do not cover."""
    notes = []
    rate = res.driver_defaults.get("INT_RATE")
    rate = rate[0] if isinstance(rate, list) and rate else rate
    if isinstance(rate, float) and rate < 0:
        notes.append(f"default interest rate on debt is negative ({rate:.2%})")
    latest = res.fy_labels[-1]
    for concept in ("InterestExpense", "PretaxIncome"):
        cv = res.mapping.get(concept, {}).get(latest)
        if cv is None or cv.status == "UNMAPPED":
            notes.append(f"{concept} UNMAPPED in {latest}")
    try:
        facts = load_companyfacts(SecClient(ua), res.cik)
    except Exception as e:  # sanity notes must never break the run
        return notes + [f"companyfacts unavailable ({type(e).__name__})"]
    for concept in SIGN_CONCEPTS:
        cv = res.mapping.get(concept, {}).get(latest)
        if cv is None or not cv.value or not (cv.tag or "").startswith("us-gaap_"):
            continue
        fact = fy_values(facts, cv.tag.removeprefix("us-gaap_"), res.currency).get(latest)
        tol = max(1.0, 0.005 * abs(cv.value))
        if fact and abs(cv.value + fact) <= tol < abs(cv.value - fact):
            notes.append(f"{concept} {latest} mapped {cv.value:,.1f} but XBRL {cv.tag} = "
                         f"{fact:,.1f} (negated label carried into the model)")
    return notes


def run_one(ticker: str, ua: str, years: int, out_dir: Path, min_years: int,
            verbose: bool, strict: bool = False) -> Row:
    row = Row(ticker)
    progress = (lambda m: print(f"    [{ticker}] {m}", flush=True)) if verbose else (lambda m: None)
    t0 = time.perf_counter()
    try:
        res = build(ticker, ua, n_years=years, progress=progress)
    except UserFacingError as e:
        row.secs = time.perf_counter() - t0
        row.note = str(e)
        if ticker not in EXPECT_40F:
            row.failures.append(f"UserFacingError: {e}")
        elif "40-F" not in str(e):
            row.failures.append(f"expected a 40-F message, got: {e}")
        return row
    except Exception as e:  # build() should never let one through
        row.secs = time.perf_counter() - t0
        row.failures.append(f"unexpected {type(e).__name__}: {e}")
        return row
    row.secs = time.perf_counter() - t0
    if ticker in EXPECT_40F:
        row.failures.append("expected a 40-F UserFacingError, but the build succeeded")

    fys = res.fy_labels
    row.n_years, row.currency = len(fys), res.currency
    row.years = f"{fys[0]}-{fys[-1]}" if fys else "-"
    missing = [f"{fy}:{st}" for fy in fys for st in STYPES
               if fy not in res.statements[st].sources]
    full = sum(all(fy in res.statements[st].sources for st in STYPES) for fy in fys)
    row.stmts = f"{full}/{len(fys)}"
    row.counts = summarize(res.checks)
    row.coverage = list(res.coverage)
    row.bal = res.bal_check_max
    if res.dropped_years:
        row.note = "dropped " + ",".join(res.dropped_years)

    if len(fys) < min_years:
        row.failures.append(f"only {len(fys)} fiscal years (< {min_years})")
    if missing:
        row.failures.append("missing statements: " + ", ".join(missing))
    if row.counts.get("FAIL", 0):
        fails = [f"{c.name} {c.fy}" for c in res.checks if c.status == "FAIL"]
        row.failures.append(f"{len(fails)} FAIL checks: " + "; ".join(fails[:8])
                            + (" ..." if len(fails) > 8 else ""))
    if row.bal is None or not row.bal < BAL_TOL:
        row.failures.append(f"max |BAL_CHECK| = {row.bal} (>= {BAL_TOL})")
    row.sanity = sanity_notes(res, ua)
    if strict:
        row.failures.extend(f"sanity: {n}" for n in row.sanity)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{ticker}.xlsx"
    path.write_bytes(res.xlsx)
    row.kb = len(res.xlsx) // 1024
    refs, load_err = scan_ref_errors(path)
    if load_err:
        row.failures.append(f"openpyxl.load_workbook failed: {load_err}")
    if refs:
        row.failures.append(f"#REF! in {len(refs)} cells: " + ", ".join(refs[:8]))
    return row


def line(r: Row) -> str:
    if not r.years:
        return f"{r.ticker:<5} {r.status:<4} {r.secs:6.1f}s  error: {r.note}"
    c = r.counts
    bal = "n/a" if r.bal is None else f"{r.bal:.1e}"
    return (f"{r.ticker:<5} {r.status:<4} {r.n_years:>2}y {r.years:<13} {r.currency:<3} "
            f"stmts {r.stmts:<5} P/W/F/S {c.get('PASS', 0)}/{c.get('WARN', 0)}/"
            f"{c.get('FAIL', 0)}/{c.get('SKIP', 0)}  cov {len(r.coverage)}  bal {bal}  "
            f"{r.secs:6.1f}s  {r.kb} KB  sanity {len(r.sanity)}" + (f"  ({r.note})" if r.note else ""))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("tickers", nargs="+")
    ap.add_argument("--ua", required=True, help='SEC User-Agent, "Name email@domain"')
    ap.add_argument("--years", type=int, default=10, help="fiscal years to request")
    ap.add_argument("--min-years", type=int, default=5, help="fewest years that pass")
    ap.add_argument("--out", type=Path, default=ROOT / "out")
    ap.add_argument("-v", "--verbose", action="store_true", help="echo pipeline progress")
    ap.add_argument("--strict", action="store_true", help="treat SANITY notes as assertions")
    args = ap.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    logging.basicConfig(level=logging.WARNING, stream=sys.stderr)

    rows = []
    for t in [t.strip().upper() for t in args.tickers]:
        r = run_one(t, args.ua, args.years, args.out, args.min_years, args.verbose, args.strict)
        rows.append(r)
        print(line(r), flush=True)
        for w in r.coverage:
            print(f"      coverage: {w}", flush=True)
        for n in r.sanity:
            print(f"      SANITY: {n}", flush=True)
        for f in r.failures:
            print(f"      ASSERT: {f}", flush=True)

    print("\n=== Summary ===")
    for r in rows:
        print(line(r))
    bad = [r for r in rows if r.failures]
    print(f"\n{len(rows) - len(bad)}/{len(rows)} tickers passed"
          + (f"; failed: {', '.join(r.ticker for r in bad)}" if bad else ""))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
