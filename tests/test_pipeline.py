"""Offline end-to-end build: the OKE fixture filing served through a fake SEC client."""
from io import BytesIO
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook

from sec_model.pipeline import BuildResult, build
from sec_model.types import UserFacingError

FIXTURE = Path(__file__).parent / "fixtures" / "OKE"
CIK = 1039684
ACCESSION = "0001039684-26-000012"
FOLDER = f"https://www.sec.gov/Archives/edgar/data/{CIK}/{ACCESSION.replace('-', '')}/"
TICKERS_URL = "https://www.sec.gov/files/company_tickers.json"
EXCHANGE_URL = "https://www.sec.gov/files/company_tickers_exchange.json"
SUBMISSIONS_URL = f"https://data.sec.gov/submissions/CIK{CIK:010d}.json"
FACTS_URL = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{CIK:010d}.json"

SHEET_ORDER = ["Cover", "Checks", "Model", "Asset Inputs", "Drivers", "WACC", "DCF", "Mapping",
               "IS (As Reported)", "BS (As Reported)", "CF (As Reported)", "Sources"]


def _submissions(form="10-K", prior_10k=False):
    rows = [(ACCESSION, "2026-02-24", "2025-12-31", form, "oke-20251231.htm")]
    if prior_10k:   # a FY2024 10-K whose folder serves nothing (no FilingSummary)
        rows.insert(0, ("0001039684-25-000001", "2025-02-25", "2024-12-31", "10-K",
                        "oke-20241231.htm"))
    keys = ["accessionNumber", "filingDate", "reportDate", "form", "primaryDocument"]
    return {"filings": {"recent": {k: [r[i] for r in rows] for i, k in enumerate(keys)},
                        "files": []}}


class FakeClient:
    """Serves the OKE fixture for the URLs the pipeline requests; 404s everything else."""

    def __init__(self, form="10-K", prior_10k=False):
        self.json = {
            TICKERS_URL: {"0": {"cik_str": CIK, "ticker": "OKE", "title": "ONEOK INC /NEW/"}},
            EXCHANGE_URL: {"fields": ["cik", "name", "ticker", "exchange"], "data": []},
            SUBMISSIONS_URL: _submissions(form, prior_10k),
            FACTS_URL: {"facts": {"us-gaap": {}, "dei": {}}},
        }
        self.text = {
            FOLDER + "FilingSummary.xml": (FIXTURE / "FilingSummary.xml").read_text(encoding="utf-8"),
            FOLDER + "R3.htm": (FIXTURE / "IS.htm").read_text(encoding="utf-8"),
            FOLDER + "R7.htm": (FIXTURE / "BS.htm").read_text(encoding="utf-8"),
            FOLDER + "R9.htm": (FIXTURE / "CF.htm").read_text(encoding="utf-8"),
        }

    def get_json(self, url: str) -> dict:
        if url not in self.json:
            raise requests.HTTPError(f"404 for {url}")
        return self.json[url]

    def get_text(self, url: str) -> str:
        if url not in self.text:
            raise requests.HTTPError(f"404 for {url}")
        return self.text[url]


@pytest.fixture(scope="module")
def oke() -> BuildResult:
    messages: list[str] = []
    r = build("OKE", "test admin@example.com", client=FakeClient(), progress=messages.append)
    assert messages == r.log
    return r


def test_build_offline_fixture_ticker_produces_workbook(oke):
    assert oke.company == "ONEOK INC /NEW/"
    assert (oke.ticker, oke.cik, oke.currency) == ("OKE", CIK, "USD")
    assert oke.fy_labels == ["FY2024", "FY2025"]
    assert oke.fcst_labels == ["FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"]
    assert oke.dropped_years == ["FY2023"]
    assert oke.company_fye == "December 31"
    assert set(oke.statements) == {"IS", "BS", "CF"}
    assert oke.checks
    assert not [c for c in oke.checks if c.status == "FAIL"]
    assert len(oke.xlsx) > 0
    wb = load_workbook(BytesIO(oke.xlsx))
    assert wb.sheetnames == SHEET_ORDER
    assert wb.active.title == "Cover"
    assert any("no filing with all three statements; dropped" in m for m in oke.log)
    assert all(m.isascii() for m in oke.log)   # default progress=print on a cp1252 console


def test_forecast_balances_in_result(oke):
    assert oke.da_basis == "ppe"
    assert oke.driver_defaults and oke.scalar_defaults
    assert set(oke.fcst_labels) <= set(oke.forecast["BAL_CHECK"])
    assert oke.bal_check_max < 1e-6


def test_year_without_statements_covered_by_later_comparatives():
    messages: list[str] = []
    r = build("OKE", "test admin@example.com", client=FakeClient(prior_10k=True),
              progress=messages.append)
    assert r.fy_labels == ["FY2024", "FY2025"]
    assert r.dropped_years == ["FY2023"]
    assert ("FY2024: no filing with all three statements; "
            "covered by a later filing's comparative columns") in messages
    assert r.statements["BS"].sources["FY2024"].accession == ACCESSION


def test_build_unknown_ticker_user_facing():
    with pytest.raises(UserFacingError, match="ZZZQ"):
        build("ZZZQ", "test admin@example.com", client=FakeClient(), progress=lambda m: None)


def test_build_40f_only_filer_user_facing():
    with pytest.raises(UserFacingError) as exc:
        build("OKE", "test admin@example.com", client=FakeClient(form="40-F"),
              progress=lambda m: None)
    assert "OKE" in str(exc.value) and "40-F" in str(exc.value)
    assert "CIK" not in str(exc.value)


def test_unexpected_error_wrapped_as_user_facing(monkeypatch, caplog):
    import sec_model.pipeline as pipeline

    def boom(*a, **k):
        raise KeyError("x")

    monkeypatch.setattr(pipeline, "resolve", boom)
    with pytest.raises(UserFacingError, match=r"Build failed for OKE: KeyError\. See server log\."):
        build("oke", "test admin@example.com", client=FakeClient(), progress=lambda m: None)
    assert any(rec.levelname == "ERROR" and rec.exc_info for rec in caplog.records)


def test_api_cross_check_reads_each_year_with_that_years_tag():
    # TRGP: FY2016 revenue is mapped via RevenueFromContractWithCustomer... (6,690.9) and
    # later years via Revenues; companyfacts' FY2016 "Revenues" is a different figure.
    from sec_model.pipeline import _facts_values
    from sec_model.types import ConceptValue

    def entry(year, val):
        return {"start": f"{year}-01-01", "end": f"{year}-12-31", "val": val * 1e6,
                "fp": "FY", "form": "10-K", "filed": f"{year + 1}-02-20"}

    facts = {"facts": {"us-gaap": {
        "Revenues": {"units": {"USD": [entry(2016, 5586.7), entry(2017, 8814.9)]}},
        "RevenueFromContractWithCustomerExcludingAssessedTax":
            {"units": {"USD": [entry(2016, 6690.9), entry(2017, 7800.8)]}},
    }}}
    rev = {"FY2016": ConceptValue("Revenue", "FY2016", 6690.9, "k",
                                  "us-gaap_RevenueFromContractWithCustomerExcludingAssessedTax",
                                  "Total revenues", "OK"),
           "FY2017": ConceptValue("Revenue", "FY2017", 8814.9, "k", "us-gaap_Revenues",
                                  "Total revenues", "OK")}
    out = _facts_values({"Revenue": rev}, facts, "USD")
    assert out["Revenue"] == pytest.approx({"FY2016": 6690.9, "FY2017": 8814.9})
