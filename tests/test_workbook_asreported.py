from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from sec_model.concepts import CONCEPTS, Concept
from sec_model.types import (CheckResult, ConceptValue, MergedRow, MergedStatement,
                             SourceRec)
from sec_model.workbook import styles as S
from sec_model.workbook.asreported import SHEET_NAMES, unit_note, write_as_reported
from sec_model.workbook.inputs_sheets import unit_for
from sec_model.workbook.mapping_sheet import write_mapping
from sec_model.workbook.meta import CoverInfo, write_checks, write_cover, write_sources

FYS = ["FY2023", "FY2024"]


def _src(fy):
    return SourceRec(f"0001-24-{fy[-2:]}", "10-K", date(2024, 2, 20), f"https://sec.gov/R2_{fy}.htm",
                     "In Millions")


def _row(key, label, vals, indent=0, abstract=False, unit="money", still=True):
    return MergedRow(key, "us-gaap_" + key, label, indent, abstract, unit,
                     dict(zip(FYS, vals)), still)


def _is():
    rows = [
        _row("Hdr", "Revenues [Abstract]", [None, None], abstract=True),
        _row("Revenues", "Revenues", [100.0, 120.0], indent=1),
        _row("Legacy", "Legacy item", [5.0, None], indent=1, still=False),
        _row("Total", "Total costs and expenses", [60.0, None], indent=1),
        _row("EPS", "Earnings per unit", [1.5, 1.75], unit="perunit"),
        _row("Units", "Units outstanding", [10.0, 11.0], unit="shares"),
    ]
    return MergedStatement("IS", "USD", FYS, {fy: date(int(fy[2:]), 12, 31) for fy in FYS}, rows,
                           {fy: _src(fy) for fy in FYS})


def _cv(concept, fy, value, key, status="OK", tag=None):
    return ConceptValue(concept, fy, value, key, tag or ("us-gaap_" + key if key else None),
                        "lbl", status)


def test_as_reported_layout_and_hidden_key_column():
    wb = Workbook()
    rows = write_as_reported(wb, _is(), SHEET_NAMES["IS"], "partnership")
    buf = BytesIO()
    wb.save(buf)
    ws = load_workbook(BytesIO(buf.getvalue()))["IS (As Reported)"]
    assert rows["Revenues"] == 6 and rows["Hdr"] == 5
    assert ws.column_dimensions["A"].hidden is True
    assert ws["A6"].value == "Revenues" and ws["B6"].value == "Revenues"
    assert ws["B6"].alignment.indent == 1
    assert ws["C6"].value == "$MM" and ws["C9"].value == "per unit" and ws["C10"].value == "MM units"
    assert ws["D6"].value == 100.0 and ws["E6"].value == 120.0
    assert ws["D6"].font.color.rgb.endswith("0000FF") and ws["D6"].font.name == "Calibri"
    assert ws["D6"].number_format == S.FMT_MONEY and ws["D9"].number_format == S.FMT_PERUNIT
    assert ws["D4"].value == "FY2023" and ws["E4"].value == "FY2024"
    assert ws.freeze_panes == "D5"
    assert ws.column_dimensions["B"].width == 48 and ws.column_dimensions["D"].width == 13
    assert ws["B2"].fill.start_color.rgb.endswith("1F3864") and ws["B2"].font.size == 14
    # None -> blank, never 0; abstract row has no values
    assert ws["E8"].value is None and ws["E7"].value is None and ws["D5"].value is None


@pytest.mark.parametrize("entity_type, shares, per, noun", [
    ("corporation", "MM shares", "per share", "share"),
    ("partnership", "MM units", "per unit", "unit"),
])
def test_as_reported_units_follow_entity_type(entity_type, shares, per, noun):
    wb = Workbook()
    write_as_reported(wb, _is(), SHEET_NAMES["IS"], entity_type)
    ws = wb["IS (As Reported)"]
    assert (ws["C9"].value, ws["C10"].value) == (per, shares)
    assert ws["B3"].value == f"USD in millions except per-{noun}; blue = SEC-reported value"
    assert unit_note("shares", "USD", entity_type) == shares
    assert unit_note("perunit", "USD", entity_type) == per
    # the Model's unit column uses the same helper
    assert (unit_for("units", "USD", entity_type), unit_for("perunit", "USD", entity_type)) == \
        (shares, per)
    # the default is a corporation, as for build_spec
    write_as_reported(wb2 := Workbook(), _is(), SHEET_NAMES["IS"])
    assert wb2["IS (As Reported)"]["C10"].value == "MM shares"


def test_non_usd_unit_note():
    wb = Workbook()
    stmt = _is()
    stmt.currency = "CAD"
    write_as_reported(wb, stmt, SHEET_NAMES["IS"])
    assert wb["IS (As Reported)"]["C6"].value == "CAD MM"


def test_dropped_row_label_and_style():
    wb = Workbook()
    write_as_reported(wb, _is(), SHEET_NAMES["IS"])
    ws = wb["IS (As Reported)"]
    assert ws["B7"].value == "Legacy item (no longer reported)"
    assert ws["B7"].font.italic and ws["B7"].font.color.rgb.endswith("808080")
    assert ws["D7"].font.italic
    assert ws["B5"].font.bold and ws["B5"].fill.start_color.rgb.endswith("F2F2F2")
    assert ws["B8"].font.bold and ws["B8"].border.top.style == "thin"


def test_mapping_value_formula_points_at_key():
    wb = Workbook()
    stmt = _is()
    write_as_reported(wb, stmt, SHEET_NAMES["IS"])
    mapping = {c.name: {} for c in CONCEPTS}
    mapping["Revenue"] = {fy: _cv("Revenue", fy, v, "Revenues") for fy, v in zip(FYS, (100.0, 120.0))}
    mapping["Gains"] = {fy: _cv("Gains", fy, None, None, "UNMAPPED") for fy in FYS}
    mapping["OperatingIncome"] = {
        fy: _cv("OperatingIncome", fy, 7.0, None, "DERIVED", tag="SUM(a,b)") for fy in FYS}
    info = write_mapping(wb, mapping, dict(SHEET_NAMES),
                         {"IS": stmt, "BS": stmt, "CF": stmt}, FYS)
    ws = wb["Mapping"]
    r = info["rows"]["Revenue"]
    assert r == 5 + [c.name for c in CONCEPTS].index("Revenue")
    vc, kc = info["value_cols"]["FY2024"], info["key_cols"]["FY2024"]
    f = ws.cell(r, vc).value
    keyref = ws.cell(r, kc).coordinate
    assert "MATCH(" in f and keyref in f and "'IS (As Reported)'!E:E" in f and "$A:$A" in f
    idx = "INDEX('IS (As Reported)'!E:E, MATCH(" + keyref + ", 'IS (As Reported)'!$A:$A, 0))"
    assert f == f'=IFERROR(IF({idx}="","",{idx}),"")'   # blank source stays "", never 0
    assert ws.cell(r, kc).value == "Revenues" and ws.cell(r, kc).fill.start_color.rgb.endswith("FFF2CC")
    assert ws.cell(r, 1).value == "Revenue" and ws.cell(r, 2).value == "OK"
    assert ws.cell(r, 3).value == "IS (As Reported)"
    assert info["key_cols"]["FY2023"] == info["value_cols"]["FY2024"] + 2   # one spacer column
    # unmapped: blank yellow key, still a formula; derived: hard-coded number + tag text
    g = info["rows"]["Gains"]
    assert ws.cell(g, info["key_cols"]["FY2024"]).value is None
    assert ws.cell(g, info["value_cols"]["FY2024"]).value.startswith("=IFERROR(IF(INDEX(")
    d = info["rows"]["OperatingIncome"]
    assert ws.cell(d, info["value_cols"]["FY2024"]).value == 7.0
    assert ws.cell(d, info["value_cols"]["FY2024"]).font.color.rgb.endswith("0000FF")
    assert ws.cell(d, info["key_cols"]["FY2024"]).value == "SUM(a,b)"


def test_cover_lists_filing_per_year_and_check_counts():
    wb = Workbook()
    names = ["Cover", "Checks", "Sources"]
    info = CoverInfo("Acme Corp", "ACM", 123, "December 31", "USD", FYS, ["FY2025E"],
                     date(2026, 10, 22), sheet_names=names,
                     check_counts={"PASS": 5, "WARN": 1, "FAIL": 2, "SKIP": 0},
                     filings=[("FY2023", "10-K", "0001-24-23", date(2024, 2, 20)),
                              ("FY2024", "10-K/A", "0001-25-24", date(2025, 2, 20))])
    write_cover(wb, info)
    write_checks(wb, [])
    write_sources(wb, [])
    ws = wb["Cover"]
    assert wb.sheetnames[0] == "Cover" and "Sheet" not in wb.sheetnames
    assert ws["B2"].value == "Acme Corp  |  SEC 3-Statement Model"
    cells = {c.value for row in ws.iter_rows() for c in row if c.value is not None}
    assert {"FY2023", "10-K/A", "0001-25-24", "2025-02-20", "Acme Corp", "ACM", 123} <= cells
    counts = {ws.cell(c.row, 2).value: c.value for row in ws.iter_rows() for c in row
              if c.column == 3 and ws.cell(c.row, 2).value in ("PASS", "WARN", "FAIL", "SKIP")
              and isinstance(c.value, int)}
    assert counts == {"PASS": 5, "WARN": 1, "FAIL": 2, "SKIP": 0}
    fail_lines = [c for row in ws.iter_rows() for c in row
                  if c.column == 2 and c.value == "FAIL" and c.font.bold
                  and c.font.color.rgb.endswith("C00000")]
    assert fail_lines
    links = [c for row in ws.iter_rows() for c in row if c.hyperlink is not None]
    assert [c.value for c in links] == names
    assert links[1].hyperlink.location == "'Checks'!A1"


@pytest.mark.parametrize("entity_type, text", [
    (None, "Corporation"),
    ("corporation", "Corporation"),
    ("partnership", "Partnership (units, distributions)"),
])
def test_cover_shows_entity_type_after_currency(entity_type, text):
    wb = Workbook()
    kw = {} if entity_type is None else {"entity_type": entity_type}
    write_cover(wb, CoverInfo("Acme LP", "ACM", 123, "December 31", "USD", FYS, **kw))
    ws = wb["Cover"]
    keys = [ws.cell(r, 2).value for r in range(4, 9)]
    assert keys == ["Company", "Ticker", "CIK", "Fiscal year end", "Currency"]
    assert (ws["B9"].value, ws["C9"].value) == ("Entity type", text)
    assert ws["B10"].value == "Fiscal years pulled"


def test_workbook_reopens_without_error():
    wb = Workbook()
    stmt = _is()
    write_as_reported(wb, stmt, SHEET_NAMES["IS"])
    write_checks(wb, [CheckResult("BS balances", "FY2024", "PASS", 10.0, 10.0, 0.5, ""),
                      CheckResult("x", "FY2024", "FAIL", 1.0, None, 0.5, "bad"),
                      CheckResult("y", "FY2024", "SKIP", None, None, 0.0, "no data")])
    write_sources(wb, [stmt])
    buf = BytesIO()
    wb.save(buf)
    wb2 = load_workbook(BytesIO(buf.getvalue()))
    ck = wb2["Checks"]
    assert ck.freeze_panes == "A2" and ck.auto_filter.ref == "A1:G4"
    assert ck["C2"].fill.start_color.rgb.endswith("D9EAD3")
    assert ck["C3"].fill.start_color.rgb.endswith("F4CCCC")
    assert ck["C4"].fill.start_color.rgb.endswith("EEEEEE")
    assert ck["D2"].number_format == S.FMT_MONEY and ck["E4"].value is None
    sc = wb2["Sources"]
    assert sc.max_row == 3 and sc["A2"].value == "Income Statement" and sc["B3"].value == "FY2024"
    assert sc["F2"].hyperlink.target == "https://sec.gov/R2_FY2023.htm"
    assert sc["G2"].value == "In Millions"


def _stmt(stype, keys):
    rows = [_row(k, k, [1.0, 2.0]) for k in keys]
    return MergedStatement(stype, "USD", FYS, {fy: date(int(fy[2:]), 12, 31) for fy in FYS}, rows,
                           {fy: _src(fy) for fy in FYS})


def test_mapping_sheet_follows_concept_stypes_order_for_shared_keys():
    wb = Workbook()
    statements = {"IS": _stmt("IS", ["DA", "Revenues"]), "BS": _stmt("BS", ["Cash"]),
                  "CF": _stmt("CF", ["DA", "Capex"])}
    for st, sheet in SHEET_NAMES.items():
        write_as_reported(wb, statements[st], sheet)
    da = next(c for c in CONCEPTS if c.name == "DA")
    assert da.stypes == ("IS", "CF")
    # concept resolved on the CF row, but "DA" also exists on the IS sheet
    mapping = {"DA": {fy: _cv("DA", fy, 2.0, "DA") for fy in FYS}}
    # stypes order picks IS first when both contain the key; make the CF row the only match
    statements["IS"].rows = [r for r in statements["IS"].rows if r.key != "DA"]
    info = write_mapping(wb, mapping, dict(SHEET_NAMES), statements, FYS)
    ws = wb["Mapping"]
    r = info["rows"]["DA"]
    assert "'CF (As Reported)'!" in ws.cell(r, info["value_cols"]["FY2024"]).value
    assert ws.cell(r, 3).value == "CF (As Reported)"
    # a key present on both IS and CF resolves to the first of the concept's stypes
    other = next(c for c in CONCEPTS if c.name == "Revenue")
    statements["CF"].rows.append(_row("Revenues", "Revenues", [1.0, 2.0]))
    mapping = {"Revenue": {fy: _cv("Revenue", fy, 2.0, "Revenues") for fy in FYS}}
    wb2 = Workbook()
    for st, sheet in SHEET_NAMES.items():
        write_as_reported(wb2, statements[st], sheet)
    info = write_mapping(wb2, mapping, dict(SHEET_NAMES), statements, FYS)
    assert other.stypes == ("IS",)
    assert "'IS (As Reported)'!" in wb2["Mapping"].cell(
        info["rows"]["Revenue"], info["value_cols"]["FY2024"]).value


def test_mapping_sheet_negative_sign_inside_blank_guard():
    wb = Workbook()
    statements = {"IS": _stmt("IS", ["X"]), "BS": _stmt("BS", ["Cash"]), "CF": _stmt("CF", ["Capex"])}
    for st, sheet in SHEET_NAMES.items():
        write_as_reported(wb, statements[st], sheet)
    neg = Concept("NegCapex", ("CF",), ("Capex",), None, sign=-1)
    mapping = {neg.name: {fy: _cv(neg.name, fy, -2.0, "Capex") for fy in FYS}}
    info = write_mapping(wb, mapping, dict(SHEET_NAMES), statements, FYS, concepts=[neg])
    f = wb["Mapping"].cell(info["rows"][neg.name], info["value_cols"]["FY2024"]).value
    assert f.startswith("=IFERROR(IF(INDEX(") and ',"",-INDEX(' in f


def test_mapping_sheet_tolerates_empty_years_and_missing_statements():
    wb = Workbook()
    info = write_mapping(wb, {}, {}, {}, [])
    assert len(info["rows"]) == len(CONCEPTS) and info["value_cols"] == {} and info["key_cols"] == {}
    ws = wb["Mapping"]
    assert ws["B5"].value == "UNMAPPED"

    wb = Workbook()
    stmt = _is()
    write_as_reported(wb, stmt, SHEET_NAMES["IS"])
    mapping = {c.name: {fy: _cv(c.name, fy, 1.0, "Revenues") for fy in FYS} for c in CONCEPTS}
    info = write_mapping(wb, mapping, {"IS": SHEET_NAMES["IS"]}, {"IS": stmt}, FYS)
    ws = wb["Mapping"]
    bs_concept = next(c for c in CONCEPTS if c.stypes[0] == "BS")
    r = info["rows"][bs_concept.name]
    assert ws.cell(r, 2).value == "UNMAPPED" and ws.cell(r, 3).value in ("", None)
    for fy in FYS:
        assert ws.cell(r, info["value_cols"][fy]).value is None
    rev = info["rows"]["Revenue"]
    assert str(ws.cell(rev, info["value_cols"]["FY2024"]).value).startswith("=IFERROR(")


def test_mapping_sheet_shared_key_uses_concept_scope_not_dict_order():
    wb = Workbook()
    statements = {"IS": _stmt("IS", ["DA"]), "BS": _stmt("BS", ["Cash"]), "CF": _stmt("CF", ["DA"])}
    for st, sheet in SHEET_NAMES.items():
        write_as_reported(wb, statements[st], sheet)
    cf_first = Concept("CfFirst", ("CF", "IS"), ("DA",), None)
    mapping = {"CfFirst": {fy: _cv("CfFirst", fy, 2.0, "DA") for fy in FYS}}
    info = write_mapping(wb, mapping, dict(SHEET_NAMES), statements, FYS, concepts=[cf_first])
    ws = wb["Mapping"]
    assert "'CF (As Reported)'!" in ws.cell(5, info["value_cols"]["FY2024"]).value
    assert ws.cell(5, 3).value == "CF (As Reported)"


def test_mapping_sheet_applies_per_year_cell_sign():
    # A negated-label year (cv.sign == -1) gets "-INDEX(...)"; a normal year does not.
    from dataclasses import replace
    wb = Workbook()
    statements = {"IS": _stmt("IS", ["X"]), "BS": _stmt("BS", ["Cash"]), "CF": _stmt("CF", ["Capex"])}
    for st, sheet in SHEET_NAMES.items():
        write_as_reported(wb, statements[st], sheet)
    c = Concept("Interest", ("IS",), ("X",), None)
    mapping = {c.name: {"FY2023": _cv(c.name, "FY2023", 2.0, "X"),
                        "FY2024": replace(_cv(c.name, "FY2024", 3.0, "X"), sign=-1)}}
    info = write_mapping(wb, mapping, dict(SHEET_NAMES), statements, FYS, concepts=[c])
    ws, r = wb["Mapping"], info["rows"][c.name]
    assert ',"",INDEX(' in ws.cell(r, info["value_cols"]["FY2023"]).value
    assert ',"",-INDEX(' in ws.cell(r, info["value_cols"]["FY2024"]).value
