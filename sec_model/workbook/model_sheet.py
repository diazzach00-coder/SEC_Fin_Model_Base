"""Model sheet: the linked 3-statement model rendered from a `ModelSpec`.

History columns pull concept lines from the Mapping sheet (``=Mapping!<cell>``) and
render derived lines with `formula.to_excel`; forecast columns render each line's
``fcst`` off the Drivers and Revenue Streams sheets (written by `inputs_sheets`). The
Model, Drivers and Revenue Streams sheets share column letters per year.

Blank-safety: a Mapping cell for an unreported year holds ``""`` (text), and a Model
history cell linking to it shows blank. Excel arithmetic on ``""`` is ``#VALUE!`` and
``""=0`` is FALSE, whereas `forecast.compute` treats a missing value as 0. So every
formula reference to a concept-sourced line's HISTORY cell renders as ``N(<cell>)``.

First history column: a `blank_first` line's cell is left empty (`model_spec` docstring).
References to it from that column render as ``N(<cell>)`` too (0, as Python reads None),
and a ``Hist`` average whose range would hold only that empty cell (one history year)
renders as ``AVERAGE(0)`` rather than ``#DIV/0!``.
"""
from dataclasses import dataclass, field

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment

from sec_model.formula import RenderContext, to_excel
from sec_model.model_spec import ModelSpec
from . import styles as S
from .inputs_sheets import (
    DRIVERS_SHEET, FONT_FORMULA_BOLD, FONT_NOTE, NUMBER_FORMATS, STREAMS_SHEET, frame,
    plan_drivers, section_header, unit_for, write_drivers, write_streams,
)

MODEL_SHEET = "Model"
SECTIONS = (("INCOME STATEMENT", "IS"), ("RATIOS & DRIVER HISTORY", "RATIO"),
            ("CASH FLOW STATEMENT", "CF"), ("BALANCE SHEET", "BS"), ("BALANCE CHECK", "CHECK"))


@dataclass
class ModelLayout:
    line_row: dict[str, int]                    # line key -> Model row
    driver_row: dict[str, int]                  # driver key -> Drivers row
    scalar_cell: dict[str, str]                 # scalar key -> "Drivers!$C$<row>"
    first_hist_col: int                         # 4 (D)
    first_fcst_col: int
    last_col: int                               # last forecast column
    # ("ASSET_REV"|"ASSET_COST", fy) -> "'Revenue Streams'!G31"
    asset_total_cell: dict[tuple[str, str], str]
    asset_total_row: dict[str, int] = field(default_factory=dict)  # block -> its total row
    sheet: str = MODEL_SHEET
    drivers_sheet: str = DRIVERS_SHEET
    streams_sheet: str = STREAMS_SHEET
    entity_type: str = "corporation"        # spec.entity_type: picks the label words


def _plan_rows(spec: ModelSpec) -> tuple[dict[str, int], list[tuple[str, int]]]:
    """Model rows for every line and section header (blank row between sections)."""
    line_row, headers, r = {}, [], 5
    for i, (title, section) in enumerate(SECTIONS):
        r += 1 if i else 0
        headers.append((title, r))
        r += 1
        for ld in spec.lines:
            if ld.section == section:
                line_row[ld.key] = r
                r += 1
    missing = set(spec.line_keys()) - set(line_row)
    if missing:
        raise ValueError(f"lines in unknown sections: {sorted(missing)}")
    return line_row, headers


def render_context(layout: ModelLayout, spec: ModelSpec, col: int,
                   qualified: bool = False) -> RenderContext:
    """The context the Model writer renders column ``col`` with (see module docstring).
    ``qualified``: line cells and history ranges carry the ``Model!`` prefix, for formulas
    on other sheets (the Drivers year-1 defaults)."""
    concept_lines = {ld.key for ld in spec.lines if isinstance(ld.hist, str)}
    blank_first = {ld.key for ld in spec.lines if ld.blank_first}
    pre = f"{layout.sheet}!" if qualified else ""

    def line_cell(key: str, c: int) -> str:
        ref = f"{pre}{S.col(c)}{layout.line_row[key]}"
        blank_safe = ((key in concept_lines and c < layout.first_fcst_col)
                      or (key in blank_first and c == layout.first_hist_col))
        return f"N({ref})" if blank_safe else ref

    def hist_range(key: str, n: int) -> str:
        last = layout.first_fcst_col - 1
        first = max(layout.first_hist_col, last - n + 1)
        if key in blank_first and last == layout.first_hist_col:
            return "0"                  # only the empty first cell: AVERAGE(0) = 0
        r = layout.line_row[key]
        return f"{pre}${S.col(first)}${r}:${S.col(last)}${r}"

    return RenderContext(
        col=col,
        line_cell=line_cell,
        driver_cell=lambda key, c: f"{layout.drivers_sheet}!{S.col(c)}{layout.driver_row[key]}",
        scalar_cell=lambda key: layout.scalar_cell[key],
        block_range=lambda key, c:
            f"'{layout.streams_sheet}'!{S.col(c)}{layout.asset_total_row[key]}",
        hist_range=hist_range,
        first_col=layout.first_hist_col,
    )


def _history_value(ld, fy, mapping_rows, mapping_value_cols, ctx, first: bool):
    """(value, is_link) for a history cell; ``first``: the first history column."""
    if first and ld.blank_first:
        return None, False
    if isinstance(ld.hist, str):
        r, vc = mapping_rows.get(ld.hist), mapping_value_cols.get(fy)
        if r is None or vc is None:
            return None, False
        return f"=Mapping!{S.col(vc)}{r}", True
    if ld.hist is None:
        return None, False
    return to_excel(ld.hist, ctx), False


def _write_line(ws, ld, r, layout, ctxs, fy_labels, fcst_labels, mapping_rows,
                mapping_value_cols, currency, note_col):
    label = ws.cell(r, 2, ld.label)
    label.font = S.FONT_BOLD if ld.bold else S.FONT_BASE
    if ld.hidden:
        label.alignment = Alignment(indent=1)
        ws.row_dimensions[r].hidden = True
        ws.row_dimensions[r].outlineLevel = 1
    ws.cell(r, 3, unit_for(ld.fmt, currency, layout.entity_type)).font = S.FONT_BASE
    if ld.bold:
        label.fill = ws.cell(r, 3).fill = S.FILL_SUBTOTAL
    for i, fy in enumerate(list(fy_labels) + list(fcst_labels)):
        c = layout.first_hist_col + i
        ctx = ctxs[c]
        cell = ws.cell(r, c)
        is_link = False
        if c < layout.first_fcst_col:
            cell.value, is_link = _history_value(ld, fy, mapping_rows, mapping_value_cols,
                                                   ctx, c == layout.first_hist_col)
            if ld.bold:
                cell.fill = S.FILL_SUBTOTAL
        else:
            cell.value = None if ld.fcst is None else to_excel(ld.fcst, ctx)
            cell.fill = S.FILL_FORECAST
        if is_link:
            cell.font = S.FONT_INPUT_BOLD if ld.bold else S.FONT_INPUT
        else:
            cell.font = FONT_FORMULA_BOLD if ld.bold else S.FONT_FORMULA
        if ld.bold:
            cell.border = S.BORDER_TOP
        cell.number_format = NUMBER_FORMATS.get(ld.fmt, S.FMT_MONEY)
    if ld.note:
        note = ws.cell(r, note_col, ld.note)
        note.data_type = "s"            # a note starting with "=" is text, never a formula
        note.font = FONT_NOTE


def write_model(wb, spec: ModelSpec, mapping_rows: dict[str, int],
                mapping_value_cols: dict[str, int], fy_labels: list[str],
                fcst_labels: list[str], driver_defaults: dict[str, float],
                scalar_defaults: dict[str, float], currency: str) -> ModelLayout:
    """Write the Model, Revenue Streams and Drivers sheets; return the Model layout.

    ``mapping_rows`` / ``mapping_value_cols`` are `write_mapping`'s ``rows`` (concept ->
    row) and ``value_cols`` (fy -> column). A concept or year missing from them leaves
    the history cell blank. Driver year-1 cells and Expr scalar defaults are written as
    formulas over this sheet (`inputs_sheets` module docstring); ``driver_defaults`` /
    ``scalar_defaults`` (`forecast.driver_defaults` / `forecast.scalar_defaults`) supply
    only the constant defaults.
    """
    if not fy_labels:
        raise ValueError("at least one history year is required")
    ws = S.get_sheet(wb, MODEL_SHEET)
    wa = S.get_sheet(wb, STREAMS_SHEET)
    wd = S.get_sheet(wb, DRIVERS_SHEET)
    line_row, headers = _plan_rows(spec)
    driver_row, scalar_cell = plan_drivers(spec)
    total_row = write_streams(wa, fy_labels, fcst_labels, MODEL_SHEET, line_row, driver_row)

    first_fcst = S.DSC + len(fy_labels)
    last_col = first_fcst + len(fcst_labels) - 1
    layout = ModelLayout(
        line_row=line_row, driver_row=driver_row, scalar_cell=scalar_cell,
        first_hist_col=S.DSC, first_fcst_col=first_fcst, last_col=last_col,
        asset_total_cell={(k, fy): f"'{STREAMS_SHEET}'!{S.col(first_fcst + i)}{row}"
                          for k, row in total_row.items() for i, fy in enumerate(fcst_labels)},
        asset_total_row=total_row, entity_type=spec.entity_type,
    )
    write_drivers(wb, wd, spec, fy_labels, fcst_labels, driver_defaults, scalar_defaults,
                  currency, render_context(layout, spec, first_fcst - 1, qualified=True))

    note_col = last_col + 2
    frame(ws, "3-Statement Model",
          f"All figures in {currency} MM unless noted; blue = SEC data via Mapping; "
          "black = formula; yellow = input",
          ["Line item", "Unit"], fy_labels, fcst_labels, {note_col: "Note"})
    S.set_widths(ws, {"B": 46, "C": 10, last_col + 1: 3, note_col: 60})
    for title, r in headers:
        section_header(ws, r, title, last_col)
    ctxs = {c: render_context(layout, spec, c) for c in range(S.DSC, last_col + 1)}
    for ld in spec.lines:
        _write_line(ws, ld, line_row[ld.key], layout, ctxs, fy_labels, fcst_labels,
                    mapping_rows, mapping_value_cols, currency, note_col)

    r = line_row.get("BAL_CHECK")
    if r is not None:
        rng = f"{S.col(S.DSC)}{r}:{S.col(last_col)}{r}"
        for op, f in (("greaterThan", "0.5"), ("lessThan", "-0.5")):
            ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=[f],
                                                          fill=S.FILL_FAIL))
    r = line_row.get("LIQ_SHORTFALL")
    if r is not None:
        first, last = S.col(first_fcst), S.col(last_col)
        ws.conditional_formatting.add(f"{first}{r}:{last}{r}", CellIsRule(
            operator="greaterThan", formula=["0.5"], fill=S.FILL_FAIL))
        rev = line_row.get("REVOLVER")
        if rev is not None:     # the capped revolver turns red in a shortfall year
            ws.conditional_formatting.add(f"{first}{rev}:{last}{rev}", FormulaRule(
                formula=[f"{first}{r}>0.5"], fill=S.FILL_FAIL))
    return layout
