"""Shared workbook styles: fonts, fills, number formats and small layout helpers."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
BLUE = "0000FF"
FONT_NAME = "Calibri"

FMT_MONEY = '#,##0.0;(#,##0.0);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_UNITS = '#,##0.0;(#,##0.0);"-"'
FMT_PERUNIT = '#,##0.00;(#,##0.00);"-"'
FMT_X = '0.0"x"'
FMT_DAYS = '0.0'

DSC = 4  # first history year column (D)


def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_, end_color=hex_)


FILL_TITLE = _fill(NAVY)
FILL_HEADER = _fill("BFBFBF")
FILL_SECTION = _fill("F2F2F2")
FILL_SUBTOTAL = _fill("E8E8E8")
FILL_FORECAST = _fill("EBF3FB")
FILL_INPUT = _fill("FFF2CC")
FILL_OUTPUT = _fill("D9EAD3")
FILL_PASS = _fill("D9EAD3")
FILL_WARN = _fill("FFF2CC")
FILL_FAIL = _fill("F4CCCC")
FILL_SKIP = _fill("EEEEEE")
STATUS_FILLS = {"PASS": FILL_PASS, "WARN": FILL_WARN, "FAIL": FILL_FAIL, "SKIP": FILL_SKIP}

FONT_BASE = Font(name=FONT_NAME, size=10)
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True)
FONT_INPUT = Font(name=FONT_NAME, size=10, color=BLUE)          # hard-coded inputs
FONT_INPUT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color=BLUE)
FONT_FORMULA = Font(name=FONT_NAME, size=10, color="000000")    # formulas
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name=FONT_NAME, size=10, italic=True, color="FFFFFF")
FONT_DROPPED = Font(name=FONT_NAME, size=10, italic=True, color="808080")
FONT_LINK = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")
FONT_FAIL = Font(name=FONT_NAME, size=11, bold=True, color="C00000")

BORDER_TOP = Border(top=Side(style="thin"))


def col(n: int) -> str:
    return get_column_letter(n)


def get_sheet(wb, name: str, index: int | None = None):
    """Create sheet ``name``; drops openpyxl's pristine default 'Sheet' if present."""
    if "Sheet" in wb.sheetnames and name != "Sheet":
        default = wb["Sheet"]
        if default.max_row == 1 and default.max_column == 1 and default["A1"].value is None:
            wb.remove(default)
    return wb.create_sheet(name, index) if index is not None else wb.create_sheet(name)


def title_bar(ws, row: int, text: str, last_col: int, subtitle: str | None = None) -> None:
    """Navy bar across A..last_col; white bold 14pt text in B. Optional subtitle on row+1."""
    for r in (row, row + 1) if subtitle is not None else (row,):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = FILL_TITLE
    ws.cell(row, 2, text).font = FONT_TITLE
    ws.row_dimensions[row].height = 22
    if subtitle is not None:
        ws.cell(row + 1, 2, subtitle).font = FONT_SUBTITLE


def header_row(ws, row: int, labels: list[str], first_col: int,
               is_forecast: list[bool] | None = None) -> None:
    """Gray bold header cells; forecast columns get the forecast fill."""
    for i, text in enumerate(labels):
        c = ws.cell(row, first_col + i, text)
        c.font = FONT_BOLD
        c.fill = FILL_FORECAST if (is_forecast and is_forecast[i]) else FILL_HEADER
        c.alignment = Alignment(horizontal="center" if i or first_col > 2 else "left")


def set_widths(ws, widths: dict) -> None:
    """``widths`` keys are column letters or 1-based indexes."""
    for k, w in widths.items():
        ws.column_dimensions[k if isinstance(k, str) else col(k)].width = w
