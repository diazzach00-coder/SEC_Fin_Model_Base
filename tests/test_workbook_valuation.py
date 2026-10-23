import re
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from sec_model.forecast import driver_defaults, scalar_defaults
from sec_model.model_spec import build_spec
from sec_model.types import MarketInputs
from sec_model.workbook import styles as S
from sec_model.workbook.valuation import DcfLayout, WaccLayout, write_dcf, write_wacc
from sec_model.workbook.model_sheet import write_model
from test_forecast import FCST, FY, TOY
from test_workbook_model import _mapping_stub

LAST_FYE = date(2025, 12, 31)


def _build(price=30.0, currency="USD", basic_units=98.5):
    spec = build_spec("ppe")
    wb = Workbook()
    rows, value_cols = _mapping_stub(wb, TOY)
    layout = write_model(wb, spec, rows, value_cols, FY, FCST, driver_defaults(spec, TOY, FY),
                         scalar_defaults(spec, TOY, FY), currency)
    market = MarketInputs(price=price, beta=1.1, risk_free=0.04, mrp=0.055)
    wl = write_wacc(wb, layout, market, currency, basic_units)
    dl = write_dcf(wb, layout, wl, market, FCST, currency, LAST_FYE,
                   company="Toy Midstream LP", ticker="TOY")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf), layout, wl, dl


def _at(wb, ref):
    """Cell for an absolute ref like "WACC!$E$20"."""
    sheet, addr = ref.split("!")
    return wb[sheet.strip("'")][addr.replace("$", "")]


def _addr(ref):
    return ref.split("!")[1].replace("$", "")


def _row_of(ws, label, col=2):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise AssertionError(f"label {label!r} not found")


def _is_input(cell):
    return cell.fill.start_color.rgb.endswith("FFF2CC")


def test_wacc_formula_references_weights_and_costs():
    wb, layout, wl, _ = _build()
    assert isinstance(wl, WaccLayout)
    ws = wb["WACC"]
    assert ws["B2"].value == "WACC Build-up"
    for ref in (wl.wacc_cell, wl.tax_cell, wl.kd_cell, wl.kd_after_tax_cell, wl.ke_cell,
                wl.price_cell, wl.diluted_units_cell, wl.net_debt_cell, wl.beta_cell,
                wl.equity_weight_cell, wl.debt_weight_cell, wl.market_cap_cell):
        assert re.fullmatch(r"WACC!\$[A-Z]+\$\d+", ref), ref
    kd_at, ke, we, wd = (_addr(r) for r in (wl.kd_after_tax_cell, wl.ke_cell,
                                             wl.equity_weight_cell, wl.debt_weight_cell))
    wacc = _at(wb, wl.wacc_cell)
    assert wacc.value == f"={kd_at}*{wd}+{ke}*{we}"
    assert wacc.fill.start_color.rgb.endswith("D9EAD3") and wacc.font.bold
    # costs
    G = S.col(layout.first_fcst_col)
    kd = _at(wb, wl.kd_cell)
    assert kd.value == f"=Drivers!{G}{layout.driver_row['INT_RATE']}" and _is_input(kd)
    assert _at(wb, wl.tax_cell).value == f"=Drivers!{G}{layout.driver_row['TAX_RATE']}"
    assert _at(wb, wl.kd_after_tax_cell).value == \
        f"={_addr(wl.kd_cell)}*(1-{_addr(wl.tax_cell)})"
    r_rf, r_beta, r_mrp = (_row_of(ws, t) for t in ("Risk-free rate", "Beta", "Market risk premium"))
    assert ws.cell(r_rf, 3).value == pytest.approx(0.04) and _is_input(ws.cell(r_rf, 3))
    assert ws.cell(r_mrp, 3).value == pytest.approx(0.055)
    assert ke.endswith(str(_row_of(ws, "Cost of equity")))
    assert _at(wb, wl.ke_cell).value == f"=C{r_rf}+C{r_beta}*C{r_mrp}"
    # weights: target override wins, else market weights; debt = 1 - equity
    we_f = _at(wb, wl.equity_weight_cell).value
    mcap, nd = _addr(wl.market_cap_cell), _addr(wl.net_debt_cell)
    assert we_f.startswith("=IF(D") and f"IFERROR(C" in we_f and we_f.endswith(",0))")
    assert _at(wb, wl.debt_weight_cell).value == f"=1-{we}"
    r_eq = _row_of(ws, "Equity (market cap)")
    assert ws.cell(r_eq, 3).value == f"={mcap}" and _is_input(ws.cell(r_eq, 4))
    assert ws.cell(_row_of(ws, "Net debt"), 3).value == f"={nd}"
    # market cap is blank-safe; price yellow
    price = _addr(wl.price_cell)
    assert _at(wb, wl.market_cap_cell).value == \
        f"=IF(ISNUMBER({price}),{price}*{_addr(wl.diluted_units_cell)},0)"
    assert _at(wb, wl.price_cell).value == pytest.approx(30.0) and _is_input(_at(wb, wl.price_cell))
    # beta selector with data validation and relevered industry beta
    sel = ws.cell(_row_of(ws, "Select beta for model"), 3)
    assert sel.value == "Observed"
    dvs = [dv for dv in ws.data_validations.dataValidation if sel.coordinate in str(dv.sqref)]
    assert len(dvs) == 1 and dvs[0].type == "list" and dvs[0].formula1 == '"Observed,Industry"'
    r_obs, r_ind = _row_of(ws, "Observed beta"), _row_of(ws, "Industry beta (relevered)")
    assert ws.cell(r_obs, 3).value == pytest.approx(1.1)
    assert _at(wb, wl.beta_cell).value == \
        f'=IF({sel.coordinate}="Observed",C{r_obs},C{r_ind})'
    hdr = _row_of(ws, "Peer")
    peers = range(hdr + 1, hdr + 7)
    for r in peers:
        assert _is_input(ws.cell(r, 2)) and ws.cell(r, 2).value is None
        assert ws.cell(r, 9).value == f'=IFERROR(C{r}/(1+(1-H{r})*G{r}/F{r}),"")'
    r_avg = _row_of(ws, "Average unlevered beta")
    assert ws.cell(r_avg, 3).value == f'=IFERROR(AVERAGE(I{hdr + 1}:I{hdr + 6}),"")'
    relev = ws.cell(_row_of(ws, "Relevered to subject"), 3).value
    assert relev.startswith(f"=IFERROR(C{r_avg}*(1+(1-") and mcap in relev and nd in relev
    # net debt bridge from the Model's last history column, with overrides
    F = S.col(layout.first_fcst_col - 1)
    lr = layout.line_row
    r_debt = _row_of(ws, "Debt (incl. revolver)", 7)
    assert ws.cell(r_debt, 8).value == f"=N(Model!{F}{lr['DEBT']})+N(Model!{F}{lr['REVOLVER']})"
    assert _is_input(ws.cell(r_debt, 9))
    assert ws.cell(r_debt, 10).value == f'=IF(I{r_debt}<>"",I{r_debt},H{r_debt})'
    assert ws.cell(_row_of(ws, "Less: cash", 7), 8).value == f"=N(Model!{F}{lr['CASH']})"
    assert ws.cell(_row_of(ws, "Less: equity investments", 7), 8).value == \
        f"=N(Model!{F}{lr['EQ_INV']})"
    assert ws.cell(_row_of(ws, "Noncontrolling interests", 7), 8).value == \
        f"=N(Model!{F}{lr['NCI_BS']})"
    r_nd = int(nd[1:])
    assert ws[nd].value == f"=J{r_nd - 5}+J{r_nd - 4}+J{r_nd - 3}-J{r_nd - 2}-J{r_nd - 1}"
    # units: dei basic units value, zero-default dilution inputs, diluted = sum
    r_basic = _row_of(ws, "Basic units (MM)", 7)
    assert ws.cell(r_basic, 10).value == pytest.approx(98.5)
    du = _addr(wl.diluted_units_cell)
    assert ws[du].value == f"=SUM(J{r_basic}:J{int(du[1:]) - 1})"
    for r in range(r_basic + 1, int(du[1:])):
        assert ws.cell(r, 10).value == 0 and _is_input(ws.cell(r, 10))
    # WACC sensitivity grid: beta (columns) x MRP (rows), plain formulas
    g = _row_of(ws, "MRP \\ Beta", 3)
    assert [ws.cell(g, c).value for c in range(4, 9)] == \
        [f"=C{r_beta}-0.2", f"=C{r_beta}-0.1", f"=C{r_beta}", f"=C{r_beta}+0.1",
         f"=C{r_beta}+0.2"]
    assert [ws.cell(g + i, 3).value for i in range(1, 6)] == \
        [f"=C{r_mrp}-0.02", f"=C{r_mrp}-0.01", f"=C{r_mrp}", f"=C{r_mrp}+0.01",
         f"=C{r_mrp}+0.02"]
    assert ws.cell(g + 3, 6).value == \
        f"=${kd_at[0]}${kd_at[1:]}*${wd[0]}${wd[1:]}+($C${r_rf}+F${g}*$C{g + 3})*${we[0]}${we[1:]}"


def test_basic_units_fall_back_to_model_and_currency_note():
    wb, layout, wl, _ = _build(currency="CAD", basic_units=None)
    ws = wb["WACC"]
    F = S.col(layout.first_fcst_col - 1)
    assert ws.cell(_row_of(ws, "Basic units (MM)", 7), 10).value == \
        f"=N(Model!{F}{layout.line_row['UNITS']})"
    for name in ("WACC", "DCF"):
        note = wb[name]["B4"]
        assert note.value == "All figures in CAD MM. Enter share price in CAD."
        assert note.font.italic and note.font.color.rgb.endswith("C00000")
    usd, *_ = _build()
    assert usd["WACC"]["B4"].value is None and usd["DCF"]["B4"].value is None


def test_dcf_ufcf_rows_link_to_model():
    wb, layout, wl, dl = _build()
    assert isinstance(dl, DcfLayout)
    ws = wb["DCF"]
    lr, dr = layout.line_row, layout.driver_row
    assert ws["B2"].value == "Discounted Cash Flow Valuation"
    assert ws["D6"].value == "Toy Midstream LP" and ws["D7"].value == "TOY"
    assert ws.cell(dl.header_row, 4).value == FY[-1]
    assert [ws.cell(dl.header_row, c).value for c in range(5, 10)] == FCST
    rows = dl.row
    for i in range(len(FCST)):
        c, mc = 5 + i, S.col(layout.first_fcst_col + i)
        L = S.col(c)
        for key, line in (("REV", "REV"), ("EBITDA", "EBITDA"), ("EBIT", "EBIT"), ("DA", "DA"),
                          ("DWC", "CF_DWC"), ("CAPEX", "CAPEX")):
            assert ws.cell(rows[key], c).value == f"=Model!{mc}{lr[line]}", (key, c)
        assert ws.cell(rows["TAX_RATE"], c).value == f"=Drivers!{mc}{dr['TAX_RATE']}"
        assert ws.cell(rows["TAX"], c).value == f"={L}{rows['EBIT']}*{L}{rows['TAX_RATE']}"
        assert ws.cell(rows["NOPAT"], c).value == f"={L}{rows['EBIT']}-{L}{rows['TAX']}"
        assert ws.cell(rows["UFCF"], c).value == \
            f"={L}{rows['NOPAT']}+{L}{rows['DA']}+{L}{rows['DWC']}+{L}{rows['CAPEX']}"
        # discounting
        assert ws.cell(rows["T"], c).value == f"=({L}{rows['DISC_DATE']}-$D$10)/365"
        assert ws.cell(rows["DF"], c).value == f"=1/(1+$H$6)^{L}{rows['T']}"
        assert ws.cell(rows["PV"], c).value == f"={L}{rows['UFCF_ADJ']}*{L}{rows['DF']}"
        prev = "$D$10" if i == 0 else f"{S.col(c - 1)}{rows['FYE']}"
        assert ws.cell(rows["DISC_DATE"], c).value == \
            f"=IF($H$8=0,{L}{rows['FYE']},AVERAGE({prev},{L}{rows['FYE']}))"
        assert ws.cell(rows["FYE"], c).value == f"=EDATE({S.col(c - 1)}{rows['FYE']},12)"
    assert ws.cell(rows["UFCF_ADJ"], 5).value == f"=E{rows['UFCF']}*$D$11"
    assert ws.cell(rows["UFCF_ADJ"], 6).value == f"=F{rows['UFCF']}"
    # actual column: last history year, blank-safe where used in arithmetic
    Fm = S.col(layout.first_fcst_col - 1)
    assert ws.cell(rows["EBITDA"], 4).value == f"=Model!{Fm}{lr['EBITDA']}"
    assert ws.cell(rows["NOPAT"], 4).value == f"=N(D{rows['EBIT']})-D{rows['TAX']}"
    # header: real dates, today, stub, links to WACC
    assert isinstance(ws["D8"].value, datetime) and ws["D8"].value.date() == LAST_FYE
    assert "yy" in ws["D8"].number_format
    assert ws["D9"].value == "=EDATE(D8,12)" and ws["D10"].value == "=TODAY()"
    assert ws["D11"].value == "=IF(D10>=D9,0,MIN(1,YEARFRAC(D10,D9)))"
    assert ws["H6"].value == "=" + wl.wacc_cell
    assert ws["H8"].value == 0 and _is_input(ws["H8"])
    # terminal value, both ways; EV, equity, per unit
    I = "I"
    assert ws[dl.growth_cell.split("!")[1].replace("$", "")].value == pytest.approx(0.025)
    tv = dl.row["TV"]
    assert ws.cell(tv, 3).value == f"=C{tv - 1}/($H$6-C{tv - 2})"
    assert ws.cell(tv - 1, 3).value == f"={I}{rows['UFCF']}*(1+C{tv - 2})"
    assert ws.cell(tv, 8).value == f"=H{tv - 2}*H{tv - 1}"
    assert ws.cell(tv - 2, 8).value == f"={I}{rows['EBITDA']}"
    assert ws.cell(tv + 1, 3).value == f"=C{tv}*{I}{rows['DF']}"
    assert ws.cell(tv + 2, 3).value == f"=SUM(E{rows['PV']}:{I}{rows['PV']})"
    ev = dl.ev_cells
    assert _at(wb, ev["perpetuity"]).value == f"=C{tv + 1}+C{tv + 2}"
    assert _at(wb, ev["exit"]).value == f"=H{tv + 1}+H{tv + 2}"
    vpu = dl.value_per_unit_cells
    eq, units = dl.row["EQUITY"], dl.row["UNITS"]
    assert _at(wb, vpu["perpetuity"]).value == f'=IFERROR(C{eq}/C{units},"n/a")'
    assert _at(wb, vpu["exit"]).value == f'=IFERROR(D{eq}/D{units},"n/a")'
    assert ws.cell(units, 3).value == "=" + wl.diluted_units_cell
    assert ws.cell(dl.row["NET_DEBT"], 4).value == "=" + wl.net_debt_cell
    # multiples block: current / year 1 / terminal x perpetuity / exit
    m = dl.row["MULT_EBITDA"]
    assert [ws.cell(m, c).value for c in range(3, 9)] == [
        f'=IFERROR($C${dl.row["EV"]}/N($D${rows["EBITDA"]}),"n/a")',
        f'=IFERROR($D${dl.row["EV"]}/N($D${rows["EBITDA"]}),"n/a")',
        f'=IFERROR($C${dl.row["EV"]}/E{rows["EBITDA"]},"n/a")',
        f'=IFERROR($D${dl.row["EV"]}/E{rows["EBITDA"]},"n/a")',
        f'=IFERROR($C${dl.row["EV"]}/I{rows["EBITDA"]},"n/a")',
        f'=IFERROR($D${dl.row["EV"]}/I{rows["EBITDA"]},"n/a")']


def _all_formula_cells(wb, sheets=None):
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        for row in ws.iter_rows():
            for c in row:
                yield ws, c


def test_sensitivity_grid_cells_are_formulas_not_data_tables():
    wb, layout, wl, dl = _build()
    for ws, c in _all_formula_cells(wb):
        assert not isinstance(c.value, (DataTableFormula, ArrayFormula)), (ws.title, c.coordinate)
    ws = wb["DCF"]
    for name, grid in dl.grids.items():
        top, left = grid["corner"]
        # 5x5 grid of plain formulas, each self-contained
        for r in range(top + 1, top + 6):
            for c in range(left + 1, left + 6):
                f = ws.cell(r, c).value
                assert isinstance(f, str) and f.startswith("=IFERROR((SUMPRODUCT("), (name, r, c)
                assert f"$C{r}" in f and f"{S.col(c)}${top}" in f
        centre = ws.cell(top + 3, left + 3).value
        assert "SUMPRODUCT($E$" in centre and wl.net_debt_cell in centre
        assert _addr(dl.grid_centre_cells[name]) == f"{S.col(left + 3)}{top + 3}"
        assert ws.cell(top + 3, left).value == "=$H$6"
    top, left = dl.grids["perpetuity"]["corner"]
    gr = _addr(dl.growth_cell)
    assert [ws.cell(top, c).value for c in range(4, 9)] == \
        [f"=${gr[0]}${gr[1:]}{d}" for d in ("-0.01", "-0.005", "", "+0.005", "+0.01")]
    top, left = dl.grids["exit"]["corner"]
    mr = _addr(dl.exit_multiple_cell)
    assert [ws.cell(top, c).value for c in range(4, 9)] == \
        [f"=${mr[0]}${mr[1:]}{d}" for d in ("-1", "-0.5", "", "+0.5", "+1")]
    assert [ws.cell(r, 3).value for r in range(top + 1, top + 6)] == \
        ["=$H$6-0.01", "=$H$6-0.005", "=$H$6", "=$H$6+0.005", "=$H$6+0.01"]
    # football field: method rows read the grids, 52-week inputs are yellow
    ff = dl.football_rows
    pt, pl = dl.grids["perpetuity"]["corner"]
    rng = f"{S.col(pl + 1)}{pt + 1}:{S.col(pl + 5)}{pt + 5}"
    assert ws.cell(ff[0], 3).value == f"=MIN({rng})" and ws.cell(ff[0], 5).value == f"=MAX({rng})"
    assert ws.cell(ff[0], 4).value == f"=E{ff[0]}-C{ff[0]}"
    assert _is_input(ws.cell(ff[2], 3)) and _is_input(ws.cell(ff[2], 5))
    chart, = ws._charts
    assert (chart.type, chart.grouping) == ("bar", "stacked")


def test_exit_multiple_default_falls_back_without_price():
    wb, layout, wl, dl = _build(price=None)
    price = _at(wb, wl.price_cell)
    assert price.value is None and _is_input(price)
    ws_w = wb["WACC"]
    assert ws_w.cell(price.row, price.column + 1).value == "enter price"
    em = _at(wb, dl.exit_multiple_cell)
    assert _is_input(em)
    p = wl.price_cell
    assert em.value == (f'=IFERROR(IF({p}="",10,({wl.market_cap_cell}+{wl.net_debt_cell})'
                        f'/N($D${dl.row["EBITDA"]})),10)')
    for method, col in (("perpetuity", "C"), ("exit", "D")):
        up = wb["DCF"][f"{col}{dl.row['UPSIDE']}"].value
        assert up == (f'=IF({p}="","enter price",'
                      f'IFERROR({p}/{_addr(dl.value_per_unit_cells[method])}-1,"n/a"))')
    with_price, _, wl2, _ = _build(price=30.0)
    assert _at(with_price, wl2.price_cell).value == pytest.approx(30.0)
    assert with_price["WACC"].cell(price.row, price.column + 1).value is None


def test_no_ref_errors_anywhere():
    for kw in ({}, {"price": None, "basic_units": None, "currency": "CAD"}):
        wb, *_ = _build(**kw)
        for ws, c in _all_formula_cells(wb):
            assert "#REF!" not in str(c.value), (ws.title, c.coordinate)


# --- formula validity scan ------------------------------------------------------------

_REF = re.compile(r"(?:(?P<sheet>'[^']+'|[A-Za-z][A-Za-z ]*)!)?\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)")


def _strip_strings(f):
    return re.sub(r'"[^"]*"', '""', f)


def _iferror_arg_counts(f):
    f = _strip_strings(f)
    out = []
    for m in re.finditer(r"IFERROR\(", f):
        depth, args, i = 1, 1, m.end()
        while depth:
            ch = f[i]
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            elif ch == "," and depth == 1:
                args += 1
            i += 1
        out.append(args)
    return out


def test_formula_validity_scan():
    for kw in ({}, {"price": None, "basic_units": None}):
        wb, *_ = _build(**kw)
        for ws, c in _all_formula_cells(wb):
            assert not isinstance(c.value, (DataTableFormula, ArrayFormula))
        n = 0
        for ws, c in _all_formula_cells(wb, {"WACC", "DCF"}):
            f = c.value
            if not (isinstance(f, str) and f.startswith("=")):
                continue
            n += 1
            where = (ws.title, c.coordinate, f)
            assert f.count('"') % 2 == 0, where
            bare = _strip_strings(f)
            depth = 0
            for ch in bare:
                depth += ch == "("
                depth -= ch == ")"
                assert depth >= 0, where
            assert depth == 0, where
            assert all(k == 2 for k in _iferror_arg_counts(f)), where
            assert "#REF!" not in f and "{" not in f, where
            # every referenced cell exists: written, or a (blank) yellow input
            for m in _REF.finditer(bare):
                if bare[max(0, m.start() - 1):m.start()].isalpha():
                    continue  # part of a function name such as LOG10
                sheet = (m.group("sheet") or ws.title).strip("'")
                assert sheet in wb.sheetnames, where
                target = wb[sheet].cell(int(m.group("row")),
                                        column_index_from_string(m.group("col")))
                assert target.value is not None or _is_input(target) or sheet == "Model", \
                    (where, sheet, target.coordinate)
        assert n > 150
