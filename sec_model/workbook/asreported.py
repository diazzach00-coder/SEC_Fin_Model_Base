"""As-Reported statement sheets: every merged row, hard-coded SEC values in blue."""
from openpyxl.styles import Alignment

from sec_model.labels import label, unit_noun
from sec_model.types import MergedStatement
from . import styles as S

SHEET_NAMES = {"IS": "IS (As Reported)", "BS": "BS (As Reported)", "CF": "CF (As Reported)"}
STATEMENT_TITLES = {"IS": "Income Statement", "BS": "Balance Sheet", "CF": "Cash Flow Statement"}
_FORMATS = {"money": S.FMT_MONEY, "shares": S.FMT_UNITS, "perunit": S.FMT_PERUNIT}


def unit_note(unit: str, currency: str, entity_type: str = "corporation") -> str:
    """Unit-column text for every sheet: "money" -> "$MM" / "CAD MM"; "shares" -> "MM shares"
    / "MM units" and "perunit" -> "per share" / "per unit" by entity type (`labels`)."""
    if unit == "money":
        return "$MM" if currency == "USD" else f"{currency} MM"
    if unit == "shares":
        return f"MM {unit_noun(entity_type)}s"
    if unit == "perunit":
        return label("per_unit", entity_type)
    return ""


def write_as_reported(wb, stmt: MergedStatement, sheet_name: str,
                      entity_type: str = "corporation") -> dict[str, int]:
    """Write one statement; returns the sheet row of every ``MergedRow.key``. Share-count
    and per-share rows are labelled shares or units by ``entity_type``."""
    ws = S.get_sheet(wb, sheet_name)
    n_years = len(stmt.fy_labels)
    last_col = S.DSC + n_years - 1
    S.title_bar(ws, 2, f"{STATEMENT_TITLES.get(stmt.stype, stmt.stype)} (As Reported)", last_col,
                subtitle=f"{stmt.currency} in millions except per-{unit_noun(entity_type)}; "
                         "blue = SEC-reported value")
    S.header_row(ws, 4, ["Key", "Line item", "Unit"] + list(stmt.fy_labels), 1)
    ws.cell(4, 2).alignment = Alignment(horizontal="left")
    S.set_widths(ws, {"A": 30, "B": 48, "C": 10})
    for i in range(n_years):
        S.set_widths(ws, {S.DSC + i: 13})
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = "D5"

    rows: dict[str, int] = {}
    for r, row in enumerate(stmt.rows, start=5):
        rows[row.key] = r
        text = row.label + ("" if row.still_reported else " (no longer reported)")
        total = (not row.is_abstract) and row.label.strip().lower().startswith("total")
        base = S.FONT_BOLD if (row.is_abstract or total) else S.FONT_BASE
        ws.cell(r, 1, row.key).font = S.FONT_BASE
        lab = ws.cell(r, 2, text)
        lab.font = S.FONT_DROPPED if not row.still_reported else base
        lab.alignment = Alignment(indent=row.indent)
        ws.cell(r, 3, unit_note(row.unit, stmt.currency, entity_type)).font = S.FONT_BASE
        fmt = _FORMATS.get(row.unit, S.FMT_MONEY)
        for i, fy in enumerate(stmt.fy_labels):
            c = ws.cell(r, S.DSC + i)
            v = row.values.get(fy)
            if v is not None and not row.is_abstract:
                c.value = v
                c.number_format = fmt
            if not row.still_reported:
                c.font = S.FONT_DROPPED
            else:
                c.font = S.FONT_INPUT_BOLD if total else S.FONT_INPUT
        for c_i in range(1, last_col + 1):
            cell = ws.cell(r, c_i)
            if row.is_abstract:
                cell.fill = S.FILL_SECTION
            elif total:
                cell.fill = S.FILL_SUBTOTAL
                cell.border = S.BORDER_TOP
    return rows
