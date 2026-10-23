"""Assemble the finished workbook from the pipeline's outputs."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook

from sec_model.checks import summarize
from sec_model.model_spec import ModelSpec
from sec_model.types import CheckResult, MarketInputs, MergedStatement
from .asreported import SHEET_NAMES, write_as_reported
from .mapping_sheet import write_mapping
from .meta import CoverInfo, write_checks, write_cover, write_sources
from .model_sheet import write_model
from .valuation import write_dcf, write_wacc

SHEET_ORDER = ["Cover", "Checks", "Model", "Asset Inputs", "Drivers", "WACC", "DCF", "Mapping",
               SHEET_NAMES["IS"], SHEET_NAMES["BS"], SHEET_NAMES["CF"], "Sources"]


def fye_text(d: date) -> str:
    """``date(2025, 12, 31)`` -> ``"December 31"``."""
    return f"{d:%B} {d.day}"


def build_workbook(*, company: str, ticker: str, cik: int, currency: str,
                   fy_labels: list[str], fcst_labels: list[str], fy_ends: dict[str, date],
                   statements: dict[str, MergedStatement], mapping: dict,
                   checks: list[CheckResult], spec: ModelSpec, drivers: dict[str, float],
                   scalars: dict[str, float], market: MarketInputs,
                   basic_units: float | None, filings_per_fy: list[tuple]) -> bytes:
    """Write every sheet and return the .xlsx bytes (sheets in ``SHEET_ORDER``).

    Write order follows the dependencies: As-Reported and Mapping first (the Model links
    to Mapping rows), then Model / Asset Inputs / Drivers, WACC, DCF, Checks, Sources,
    and Cover last (it lists check counts and the sheet index).
    """
    wb = Workbook()
    stmt_sheets = {}
    for stype in ("IS", "BS", "CF"):
        if stype in statements:
            write_as_reported(wb, statements[stype], SHEET_NAMES[stype])
            stmt_sheets[stype] = SHEET_NAMES[stype]
    m = write_mapping(wb, mapping, stmt_sheets, statements, fy_labels)
    layout = write_model(wb, spec, m["rows"], m["value_cols"], fy_labels, fcst_labels,
                         drivers, scalars, currency)
    wacc = write_wacc(wb, layout, market, currency, basic_units)
    write_dcf(wb, layout, wacc, market, fcst_labels, currency, fy_ends[fy_labels[-1]],
              company=company, ticker=ticker)
    write_checks(wb, checks)
    write_sources(wb, [statements[s] for s in ("IS", "BS", "CF") if s in statements])

    index = [n for n in SHEET_ORDER if n in wb.sheetnames]   # every sheet but the Cover
    write_cover(wb, CoverInfo(
        company=company, ticker=ticker, cik=cik, fye=fye_text(fy_ends[fy_labels[-1]]),
        currency=currency, fy_labels=list(fy_labels), forecast_labels=list(fcst_labels),
        date_pulled=date.today(), sheet_names=index, check_counts=summarize(checks),
        filings=list(filings_per_fy)))

    order = [n for n in SHEET_ORDER if n in wb.sheetnames]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == order[0]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
