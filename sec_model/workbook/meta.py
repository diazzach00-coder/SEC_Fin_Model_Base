"""Cover, Checks and Sources sheets."""
from dataclasses import dataclass, field
from datetime import date

from openpyxl.styles import Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

from sec_model.types import CheckResult, MergedStatement
from . import styles as S
from .asreported import STATEMENT_TITLES

COVER, CHECKS, SOURCES = "Cover", "Checks", "Sources"


@dataclass
class CoverInfo:
    company: str
    ticker: str
    cik: int | str
    fye: str                                   # e.g. "December 31"
    currency: str
    fy_labels: list[str]                       # fiscal years pulled, oldest -> newest
    forecast_labels: list[str] = field(default_factory=list)
    date_pulled: date | None = None
    data_source_note: str = ("SEC EDGAR Financial Report (R-file) pages from each 10-K; "
                             "XBRL companyfacts API used only for cross-checks.")
    sheet_names: list[str] = field(default_factory=list)
    check_counts: dict[str, int] = field(default_factory=dict)
    filings: list[tuple] = field(default_factory=list)   # (fy, form, accession, filing_date)


def _kv(ws, r: int, key: str, value) -> None:
    ws.cell(r, 2, key).font = S.FONT_BOLD
    ws.cell(r, 3, value).font = S.FONT_BASE
    ws.cell(r, 3).alignment = Alignment(horizontal="left")


def _section(ws, r: int, text: str, last_col: int) -> None:
    for c in range(2, last_col + 1):
        ws.cell(r, c).fill = S.FILL_SECTION
    ws.cell(r, 2, text).font = S.FONT_BOLD


def write_cover(wb, info: CoverInfo):
    ws = S.get_sheet(wb, COVER, 0)
    S.title_bar(ws, 2, f"{info.company}  |  SEC 3-Statement Model", 6)
    S.set_widths(ws, {"A": 3, "B": 28, "C": 22, "D": 26, "E": 16, "F": 40})
    r = 4
    fy_rng = (f"{info.fy_labels[0]} - {info.fy_labels[-1]}" if info.fy_labels else "")
    fc_rng = (f"{info.forecast_labels[0]} - {info.forecast_labels[-1]}"
              if info.forecast_labels else "")
    for k, v in [("Company", info.company), ("Ticker", info.ticker), ("CIK", info.cik),
                 ("Fiscal year end", info.fye), ("Currency", info.currency),
                 ("Fiscal years pulled", fy_rng), ("Forecast years", fc_rng),
                 ("Date pulled", info.date_pulled.isoformat() if info.date_pulled else ""),
                 ("Data source", info.data_source_note)]:
        _kv(ws, r, k, v)
        r += 1

    r += 1
    _section(ws, r, "Filing per fiscal year", 6)
    r += 1
    S.header_row(ws, r, ["FY", "Form", "Accession", "Filing date"], 2)
    for fy, form, accession, fdate in info.filings:
        r += 1
        for i, v in enumerate([fy, form, accession,
                               fdate.isoformat() if hasattr(fdate, "isoformat") else fdate]):
            ws.cell(r, 2 + i, v).font = S.FONT_BASE
            ws.cell(r, 2 + i).alignment = Alignment(horizontal="left")

    r += 2
    _section(ws, r, "Checks", 6)
    for status in ("PASS", "WARN", "FAIL", "SKIP"):
        r += 1
        ws.cell(r, 2, status).font = S.FONT_BASE
        ws.cell(r, 2).fill = S.STATUS_FILLS[status]
        ws.cell(r, 3, info.check_counts.get(status, 0)).font = S.FONT_BASE
        ws.cell(r, 3).alignment = Alignment(horizontal="left")
    if info.check_counts.get("FAIL", 0) > 0:
        r += 1
        ws.cell(r, 2, "FAIL").font = S.FONT_FAIL
        ws.cell(r, 3, f"{info.check_counts['FAIL']} check(s) failed - see Checks sheet").font = S.FONT_FAIL

    r += 2
    _section(ws, r, "Sheet index", 6)
    for name in info.sheet_names:
        r += 1
        c = ws.cell(r, 2, name)
        c.font = S.FONT_LINK
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{name}'!A1", display=name)
    return ws


def write_checks(wb, results: list[CheckResult]):
    ws = S.get_sheet(wb, CHECKS)
    S.header_row(ws, 1, ["Check", "FY", "Status", "LHS", "RHS", "Tolerance", "Note"], 1)
    S.set_widths(ws, {"A": 38, "B": 9, "C": 9, "D": 14, "E": 14, "F": 12, "G": 70})
    for r, res in enumerate(results, start=2):
        for c, v in enumerate([res.name, res.fy, res.status, res.lhs, res.rhs,
                               res.tolerance, res.note], start=1):
            cell = ws.cell(r, c, v)
            cell.font = S.FONT_INPUT if c in (4, 5, 6) and v is not None else S.FONT_BASE
            if c in (4, 5, 6):
                cell.number_format = S.FMT_MONEY
        ws.cell(r, 3).fill = S.STATUS_FILLS.get(res.status, S.FILL_SKIP)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(len(results) + 1, 2)}"
    return ws


def write_sources(wb, statements: list[MergedStatement]):
    ws = S.get_sheet(wb, SOURCES)
    S.header_row(ws, 1, ["Statement", "FY", "Form", "Accession", "Filing date", "R-file URL",
                         "Scale text"], 1)
    S.set_widths(ws, {"A": 22, "B": 9, "C": 9, "D": 22, "E": 13, "F": 80, "G": 28})
    r = 1
    for stmt in statements:
        for fy in stmt.fy_labels:
            src = stmt.sources.get(fy)
            if src is None:
                continue
            r += 1
            fdate = src.filing_date.isoformat() if hasattr(src.filing_date, "isoformat") else src.filing_date
            for c, v in enumerate([STATEMENT_TITLES.get(stmt.stype, stmt.stype), fy, src.form,
                                   src.accession, fdate, src.r_url, src.scale_text], start=1):
                ws.cell(r, c, v).font = S.FONT_BASE
            link = ws.cell(r, 6)
            link.hyperlink = src.r_url
            link.font = S.FONT_LINK
    ws.freeze_panes = "A2"
    return ws
