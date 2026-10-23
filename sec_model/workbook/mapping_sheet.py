"""Mapping sheet: one row per concept; value cells INDEX/MATCH off an editable key block."""
from openpyxl.styles import Alignment

from sec_model.concepts import CONCEPTS
from sec_model.types import MergedStatement
from . import styles as S

SHEET_NAME = "Mapping"


def _sheet_for(cv, concept, statements: dict[str, MergedStatement]) -> str:
    """Statement type a ConceptValue resolved on (via row_key), else the concept's first scope."""
    if cv is not None and cv.row_key is not None:
        for stype, stmt in statements.items():
            if any(r.key == cv.row_key for r in stmt.rows):
                return stype
    return concept.stypes[0]


def write_mapping(wb, mapping: dict, stmt_sheets: dict[str, str],
                  statements: dict[str, MergedStatement], fy_labels: list[str],
                  concepts=None) -> dict:
    """Write the Mapping sheet.

    ``stmt_sheets``: stype -> As-Reported sheet name; ``statements``: stype -> MergedStatement
    (column of each FY on its sheet is derived from the statement's own ``fy_labels``).
    Returns {"rows": {concept: row}, "value_cols": {fy: col}, "key_cols": {fy: col}}.
    """
    concepts = CONCEPTS if concepts is None else concepts
    ws = S.get_sheet(wb, SHEET_NAME)
    n = len(fy_labels)
    value_cols = {fy: S.DSC + i for i, fy in enumerate(fy_labels)}
    spacer = S.DSC + n
    key_cols = {fy: spacer + 1 + i for i, fy in enumerate(fy_labels)}
    tag_col, label_col = spacer + 1 + n, spacer + 2 + n
    S.title_bar(ws, 2, "Concept Mapping", label_col,
                subtitle="Edit a yellow row key to repoint a concept to another As-Reported row")
    S.header_row(ws, 4, ["Concept", "Status", "Statement"], 1)
    S.header_row(ws, 4, list(fy_labels), S.DSC)
    S.header_row(ws, 4, [f"Key {fy}" for fy in fy_labels], spacer + 1)
    S.header_row(ws, 4, ["Tag (latest)", "Label (latest)"], tag_col)
    S.set_widths(ws, {"A": 26, "B": 11, "C": 18, spacer: 3, tag_col: 40, label_col: 48})
    for c in list(value_cols.values()):
        S.set_widths(ws, {c: 13})
    for c in key_cols.values():
        S.set_widths(ws, {c: 24})
    ws.freeze_panes = "D5"

    latest = fy_labels[-1]
    rows: dict[str, int] = {}
    for r, concept in enumerate(concepts, start=5):
        rows[concept.name] = r
        per_fy = mapping.get(concept.name, {})
        cv_latest = per_fy.get(latest)
        ws.cell(r, 1, concept.name).font = S.FONT_BOLD
        ws.cell(r, 2, cv_latest.status if cv_latest else "UNMAPPED").font = S.FONT_BASE
        ws.cell(r, 3, stmt_sheets[_sheet_for(cv_latest, concept, statements)]).font = S.FONT_BASE
        for fy in fy_labels:
            cv = per_fy.get(fy)
            stype = _sheet_for(cv, concept, statements)
            stmt = statements[stype]
            vcell = ws.cell(r, value_cols[fy])
            kcell = ws.cell(r, key_cols[fy])
            vcell.number_format = S.FMT_MONEY
            if cv is not None and cv.row_key is None and cv.value is not None:
                vcell.value = cv.value            # DERIVED / companyfacts: no As-Reported cell
                vcell.font = S.FONT_INPUT
                kcell.value = cv.tag or ""
                kcell.font = S.FONT_DROPPED
                continue
            kcell.fill = S.FILL_INPUT
            kcell.font = S.FONT_INPUT
            if cv is not None and cv.row_key is not None:
                kcell.value = cv.row_key
            if fy not in stmt.fy_labels:
                vcell.value = ""
                continue
            sheet = stmt_sheets[stype]
            sc = S.col(S.DSC + stmt.fy_labels.index(fy))
            sign = "-" if concept.sign == -1 else ""
            vcell.value = (f"=IFERROR({sign}INDEX('{sheet}'!{sc}:{sc}, "
                           f"MATCH({kcell.coordinate}, '{sheet}'!$A:$A, 0)), \"\")")
            vcell.font = S.FONT_FORMULA
        ws.cell(r, tag_col, (cv_latest.tag or "") if cv_latest else "").font = S.FONT_BASE
        ws.cell(r, label_col, (cv_latest.label or "") if cv_latest else "").font = S.FONT_BASE
        ws.cell(r, tag_col).alignment = Alignment(wrap_text=False)
    return {"rows": rows, "value_cols": value_cols, "key_cols": key_cols}
